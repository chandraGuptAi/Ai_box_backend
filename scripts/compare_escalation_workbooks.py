"""Compare exported Node and Python workbook cell values, without contacting any service."""

import argparse
from datetime import date, datetime
from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel


def normalized(value):
    if isinstance(value, (date, datetime)):
        return to_excel(value)
    return "" if value is None else value


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("node_workbook")
    parser.add_argument("python_workbook")
    args = parser.parse_args()
    old, new = (load_workbook(name, data_only=True) for name in (args.node_workbook, args.python_workbook))
    differences = []
    if old.sheetnames != new.sheetnames:
        differences.append("Sheet names/order differ")
    for name in old.sheetnames:
        if name not in new:
            continue
        a, b = old[name], new[name]
        for row in range(1, max(a.max_row, b.max_row) + 1):
            for col in range(1, max(a.max_column, b.max_column) + 1):
                av, bv = normalized(a.cell(row, col).value), normalized(b.cell(row, col).value)
                equal = (
                    abs(av - bv) < 1e-9 if type(av) in (float, int) and type(bv) in (float, int) else av == bv
                )
                if not equal:
                    differences.append(f"{name}!{a.cell(row,col).coordinate}: {av!r} != {bv!r}")
    print("\n".join(differences[:100]) if differences else "All cached cell values match.")
    raise SystemExit(bool(differences))


if __name__ == "__main__":
    main()
