from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app.main import app
from app.api.canonical_report_routes import engine
from app.report_engine.service import Engine
from app.report_engine.models import Operation, Chart
from app.report_engine.export import excel, pdf
from app.report_engine.registry import SkillRegistry

ROWS = [
    {'order_status': 'Delivered', 'courier_name': 'A', 'order_date': '2026-01-01'},
    {'order_status': 'Cancelled', 'courier_name': 'A', 'order_date': '2026-01-02'},
    {'order_status': 'In Transit', 'courier_name': 'B'},
    {'courier_name': 'B'},
]

@pytest.fixture
def service(tmp_path):
    return Engine(tmp_path / 'reports.db')

def test_verified_counts_and_lineage(service):
    report, _ = service.create('alice', ROWS, 'Explicit test fixture')
    assert {m.id: str(m.value) for m in report.kpis} == {
        'total':'4', 'delivered':'1', 'cancelled':'1', 'unknown':'1', 'cancellation_rate':'25.00'}
    assert report.actual_data_period == {'start':'2026-01-01', 'end':'2026-01-02'}
    assert report.source_metadata['missing_dates'] == 2
    assert all(sum(c.datasets[0].data) == 4 for c in report.charts)
    assert all(m.source_references == ['source_shipments'] for m in report.kpis)

def test_targeted_chart_changes_persist(service):
    report, token = service.create('alice', ROWS)
    untouched = report.charts[1].model_dump()
    converted = service.execute(report.report_id, 'alice', token, Operation(
        operation='UPDATE_CHART_TYPE', expected_version=1, chart_id='chart_status', chart_type='line'))
    assert converted.charts[0].type == 'line'
    assert converted.charts[0].datasets == report.charts[0].datasets
    assert converted.charts[1].model_dump() == untouched
    reloaded = Engine(service.path).get(report.report_id, 'alice', token)
    assert reloaded == converted
    removed = service.execute(report.report_id, 'alice', token, Operation(
        operation='REMOVE_CHART', expected_version=2, chart_id='chart_status'))
    assert [c.model_dump() for c in removed.charts] == [untouched]
    assert removed.kpis == report.kpis
    with pytest.raises(HTTPException) as conflict:
        service.execute(report.report_id, 'alice', token, Operation(operation='EDIT_REPORT', expected_version=1, title='stale'))
    assert conflict.value.status_code == 409

def test_concurrent_edit_is_atomic(service):
    report, token = service.create('alice', ROWS)
    def edit(title):
        try:
            return service.execute(report.report_id, 'alice', token,
                Operation(operation='EDIT_REPORT', expected_version=1, title=title)).version
        except HTTPException as exc:
            return exc.status_code
    with ThreadPoolExecutor(2) as pool:
        assert sorted(pool.map(edit, ['One', 'Two'])) == [2, 409]

@pytest.mark.parametrize('account,token_override', [('bob', None), ('alice', 'wrong')])
def test_access_is_scoped(service, account, token_override):
    report, token = service.create('alice', ROWS)
    with pytest.raises(HTTPException) as denied:
        service.get(report.report_id, account, token_override or token)
    assert denied.value.status_code == 404

def test_export_and_preview_use_same_canonical_values(service):
    report, token = service.create('alice', ROWS)
    report = service.execute(report.report_id, 'alice', token, Operation(
        operation='REMOVE_CHART', expected_version=1, chart_id='chart_courier'))
    report = service.execute(report.report_id, 'alice', token, Operation(
        operation='UPDATE_CHART_TYPE', expected_version=2, chart_id='chart_status', chart_type='line'))
    workbook = load_workbook(BytesIO(excel(report)))
    rows = list(workbook['Report'].values)
    assert [(r[0], r[1]) for r in rows[1:6]] == [(m.id, float(m.value)) for m in report.kpis]
    assert workbook.sheetnames == ['Report', 'Chart 1']
    assert workbook['Chart 1']['A1'].value == 'chart_status'
    assert workbook['Chart 1']['B1'].value == 'line'
    assert workbook['Chart 1']._charts[0].__class__.__name__ == 'LineChart'
    assert [r[1] for r in list(workbook['Chart 1'].values)[2:]] == report.charts[0].datasets[0].data
    # Inspect PDF text independently using the project's available PDF reader.
    from pypdf import PdfReader
    text = '\n'.join(page.extract_text() for page in PdfReader(BytesIO(pdf(report))).pages)
    for metric in report.kpis:
        assert metric.id in text and str(metric.value) in text
    assert 'Shipments by courier' not in text
    assert 'Shipments by status' in text
    assert 'Filters: {}' in text

def test_formula_labels_are_text(service):
    report, _ = service.create('alice', [{'courier_name': '=1+2', 'order_status':'Delivered'}])
    book = load_workbook(BytesIO(excel(report)))
    cell = book['Chart 2']['A3']
    assert cell.value == '=1+2' and cell.data_type == 's'

def test_api_flow(service):
    app.dependency_overrides[engine] = lambda: service
    try:
        client = TestClient(app)
        response = client.post('/api/canonical-reports', headers={'X-User-Id':'alice'},
            json={'records':ROWS, 'prompt':'Generate shipment analysis in graphical form'})
        assert response.status_code == 200, response.text
        created = response.json()
        report = created['report']; rid = report['report_id']
        headers = {'X-User-Id':'alice', 'Authorization':'Bearer '+created['access_token']}
        assert len(report['charts']) == 2
        assert client.get(f'/api/canonical-reports/{rid}', headers=headers).json() == report
        assert client.get(f'/api/canonical-reports/{rid}/export/xlsx?version=1', headers=headers).status_code == 200
        assert client.get(f'/api/canonical-reports/{rid}/export/pdf?version=0', headers=headers).status_code == 409
        assert client.get(f'/api/canonical-reports/{rid}', headers={'X-User-Id':'alice'}).status_code == 401
        assert client.get(f'/api/canonical-reports/{rid}', headers={**headers, 'X-User-Id':'bob'}).status_code == 404
        assert client.post(f'/api/canonical-reports/{rid}/operations', headers=headers, json={
            'operation':'ADD_CHART', 'expected_version':1, 'dimension':'state'}).status_code == 422
        assert client.post('/api/canonical-reports', headers=headers, json={
            'records':ROWS, 'prompt':'Show monthly revenue'}).status_code == 422
        assert 'database' not in str(report).lower()
    finally:
        app.dependency_overrides.pop(engine, None)

def test_registry_and_invalid_data():
    registry = SkillRegistry()
    assert [s.name for s in registry.select('REMOVE_CHART')] == ['chart-management']
    with pytest.raises(ValueError):
        registry.load('../../untrusted')
    with pytest.raises(ValueError):
        Chart(id='x', title='Invalid', labels=['a','b'], datasets=[{'label':'n','data':[1]}],
            source_references=['source'], metric_references=['total'])


def test_chart_add_and_title_update_are_targeted(service):
    report, token = service.create('alice', ROWS)
    original = [chart.model_dump() for chart in report.charts]
    added = service.execute(report.report_id, 'alice', token, Operation(
        operation='ADD_CHART', expected_version=1, dimension='status', chart_type='doughnut'))
    assert [chart.model_dump() for chart in added.charts[:2]] == original
    assert added.charts[2].id not in {chart['id'] for chart in original}
    assert added.charts[2].datasets == added.charts[0].datasets
    updated = service.execute(report.report_id, 'alice', token, Operation(
        operation='UPDATE_CHART', expected_version=2, chart_id=added.charts[2].id, title='Shipment mix'))
    assert updated.charts[2].title == 'Shipment mix'
    assert updated.charts[2].id == added.charts[2].id
    assert [chart.model_dump() for chart in updated.charts[:2]] == original


def test_zero_and_unknown_are_distinct_and_empty_source_is_rejected(service):
    report, _ = service.create('alice', [{'order_status': None}])
    assert {m.id: m.value for m in report.kpis} == {
        'total': 1, 'delivered': 0, 'cancelled': 0, 'unknown': 1, 'cancellation_rate': 0}
    with pytest.raises(HTTPException) as invalid:
        service.create('alice', [])
    assert invalid.value.status_code == 422


def test_filters_require_authorized_extraction(service):
    report, token = service.create('alice', ROWS)
    with pytest.raises(HTTPException) as unavailable:
        service.execute(report.report_id, 'alice', token,
            Operation(operation='UPDATE_FILTERS', expected_version=1))
    assert unavailable.value.status_code == 422
    assert service.get(report.report_id, 'alice', token).version == 1
