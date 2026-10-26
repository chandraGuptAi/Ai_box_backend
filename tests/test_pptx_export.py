"""PPTX exports retain report content and use existing access boundaries."""
from io import BytesIO
from pptx import Presentation
from app.models.report_schema import ReportDocument
from app.generators.pptx_generator import generate_pptx
from tests.test_reports import client, fake_ai, SAMPLE_REQUEST, STRUCTURED_AI


def test_long_content_and_tables_are_preserved():
    report = ReportDocument.model_validate({
        'report_type': 'remittance', 'report_name': 'Report',
        'generated_at': '2026-10-26T10:00:00Z',
        'content': {'title': 'Test', 'executive_summary': 'Summary',
                    'sections': [{'heading': 'Details', 'content': 'word ' * 1500 + 'ENDMARKER'}]},
        'tables': [{'title': 'Values', 'columns': ['A', 'B'],
                    'rows': [[str(i), 'value'] for i in range(35)]}]})
    deck = Presentation(BytesIO(generate_pptx(report)))
    texts = '\n'.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    assert 'ENDMARKER' in texts
    cells = [cell.text for slide in deck.slides for shape in slide.shapes if shape.has_table
             for row in shape.table.rows for cell in row.cells]
    assert '34' in cells
    assert len(deck.slides) > 5


def test_pptx_export_and_protected_download(client, fake_ai):
    fake_ai.response = STRUCTURED_AI
    generated = client.post('/api/reports/generate', json=SAMPLE_REQUEST)
    assert generated.status_code == 200
    report = generated.json()
    exported = client.post('/api/reports/export/pptx', json=report)
    assert exported.status_code == 200
    assert exported.headers['content-type'].endswith('presentationml.presentation')
    assert '.pptx' in exported.headers['content-disposition']
    deck = Presentation(BytesIO(exported.content))
    assert any(shape.has_table for slide in deck.slides for shape in slide.shapes)
    url = f"/api/reports/{report['report_id']}/download?format=pptx"
    assert client.get(url).status_code == 404
    authorized = client.get(url, headers={'Authorization': 'Bearer '+generated.headers['X-Report-Access-Token']})
    assert authorized.status_code == 200
    assert authorized.content.startswith(b'PK')


def test_saved_pptx_uses_owner_scoped_json(client, fake_ai):
    from app.api.auth import get_user_reports
    from app.main import app
    fake_ai.response = STRUCTURED_AI
    report = client.post('/api/reports/generate', json=SAMPLE_REQUEST).json()

    class OwnerReports:
        def get(self, report_id):
            return {'report': report} if report_id == report['report_id'] else None

        def file(self, *args):
            raise AssertionError('PPTX must not read an existing DOCX blob')

    app.dependency_overrides[get_user_reports] = OwnerReports
    try:
        response = client.get(f"/api/reports/saved/{report['report_id']}/download?format=pptx")
        assert response.status_code == 200
        Presentation(BytesIO(response.content))
        assert client.get('/api/reports/saved/missing/download?format=pptx').status_code == 404
    finally:
        app.dependency_overrides.pop(get_user_reports, None)


def test_preview_and_pptx_share_layout_without_one_slide_per_paragraph(client, fake_ai):
    from app.generators.presentation_layout import build_slides
    from tests.test_reports import STRUCTURED_AI
    fake_ai.response = STRUCTURED_AI
    report = client.post('/api/reports/generate', json=SAMPLE_REQUEST).json()
    exported = client.post('/api/reports/export/pptx', json=report)
    preview = client.post('/api/reports/preview/pptx', json=report)
    assert exported.status_code == preview.status_code == 200
    assert preview.headers['cache-control'] == 'no-store'
    deck = Presentation(BytesIO(exported.content))
    frames = preview.json()['slides']
    assert len(frames) == len(deck.slides)
    assert frames[0]['image'].startswith('data:image/svg+xml;base64,')
    for page in build_slides(ReportDocument.model_validate(report)):
        for item in page['items']:
            assert item['x'] >= 0 and item['y'] >= 0
            assert item['x'] + item['w'] <= 1601
            assert item['y'] + item['h'] <= 900


def test_preview_escapes_source_labels_and_does_not_execute_markup():
    import base64
    from app.generators.presentation_layout import preview_pptx
    report = ReportDocument.model_validate({'report_type': 'remittance', 'report_name': 'Report',
        'generated_at': '2026-10-26T10:00:00Z', 'content': {'title': 'Revenue & Costs', 'executive_summary': 'A & B'}})
    svg = base64.b64decode(preview_pptx(report)['slides'][0]['image'].split(',', 1)[1]).decode()
    assert 'Revenue &amp; Costs' in svg


def test_section_paragraphs_share_a_slide_with_their_evidence_table():
    from app.generators.presentation_layout import build_slides
    from tests.test_ordered_reports import generate, section
    report, _ = generate(prompt='Shipment status breakdown', sections=[section('status', ('shipment_status',))])
    layouts = build_slides(report)
    section_pages = [page for page in layouts if page['title'] == report.structured_report.sections[0].title]
    assert len(section_pages) == 1
    assert any(item['kind'] == 'table' for item in section_pages[0]['items'])
    assert any(item['kind'] == 'text' and item['y'] == 215 for item in section_pages[0]['items'])
    assert len(layouts) <= 6


def test_preview_requires_the_same_validated_shipment_revision_as_download(client, fake_ai):
    fake_ai.response = STRUCTURED_AI
    report = client.post('/api/reports/generate', json=SAMPLE_REQUEST).json()
    report['validation_token'] = None
    response = client.post('/api/reports/preview/pptx', json=report)
    assert response.status_code == 409


def test_chart_source_values_are_not_lost_when_report_has_an_unrelated_table():
    from app.generators.presentation_layout import build_slides
    from tests.test_ordered_reports import generate, section
    report, _ = generate(prompt='Courier shipment counts', sections=[section('couriers', ('courier',), charts=['bar'])])
    # A manually added chart has categories absent from the existing courier table.
    report.structured_report.charts[0].labels = ['Different category '+str(i) for i in range(len(report.structured_report.charts[0].labels))]
    layouts = build_slides(report)
    assert any(page['title'].endswith(' - Data') for page in layouts)
