"""The production plan-first workflow, over synthetic request-local shipments only."""
import asyncio
import copy
import json
from dataclasses import replace
from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfReader

from app.config import settings, Settings
from app.errors import AIResponseError, ClarificationRequiredError, ReportAccuracyError
from app.models.report_schema import ReportRequest
from app.analysis.section_reports import verify_section_coverage
from app.services.report_service import ReportService

EXAMPLE = "Create an insight report based on the given data: buyers first, top RTO pending, top courier partners, top cities with pincodes."


def shipments():
    statuses = ["Delivered", "RTO Pending", "Pickup Scheduled", "Not Picked", "Undelivered", "RTO In Transit", "RTO Delivered", "RTO"]
    return [{"order_id": f"SYNTH-{i}", "status": status, "courierName": "Courier A" if i % 3 else "Courier B",
             "consignee": {"id": "BUYER-" + str(i % 3), "name": "Synthetic Private Name",
                           "city": "City A" if i % 2 else "City B", "pincode": "110001" if i % 2 else "560001", "state": "Not a city"},
             "orderDate": "2026-10-02", "manifest_datetime": "2026-10-12T10:00:00+05:30",
             "domestic_transit_stats": {"date_pickup": "2026-10-03", "date_delivery": "2026-10-05"},
             "totalCharge": str(100+i), "invoiceTotalAmount": str(1000+i)} for i, status in enumerate(statuses)]



def section(identifier, dimensions=(), metrics=("orders",), **kwargs):
    return {"id": identifier, "title": identifier.replace("_", " ").title(),
        "requirement_ids": [identifier+"_requirement"], "dimensions": list(dimensions), "metrics": list(metrics), **kwargs}


EXAMPLE_SECTIONS = [section("recipient_overview", ("buyer",)),
    section("returns_waiting", ("courier",), filters=[{"field": "rto_stage", "op": "eq", "value": "pending"}]),
    section("partner_volume", ("courier",), ranking_metric="orders"),
    section("postal_destinations", ("destination_city", "destination_pincode"))]


class PlannedAI:
    model = "synthetic-ai"

    def __init__(self, sections=None, fail_narrative=False, alter_plan=None, events=None):
        self.sections = copy.deepcopy(EXAMPLE_SECTIONS if sections is None else sections)
        self.fail_narrative, self.alter_plan = fail_narrative, alter_plan
        self.calls, self.events = [], events if events is not None else []

    async def generate_json(self, messages, **kwargs):
        self.calls.append(messages)
        context = json.loads(messages[-1]["content"])
        if "[report_plan]" in messages[0]["content"]:
            self.events.append("plan")
            result = {"intent": "requested_shipment_analysis", "source_scope": context["source_scope"],
                "sections": self.sections, "requirement_ids": [r for s in self.sections for r in s["requirement_ids"]]}
            return self.alter_plan(result) if self.alter_plan else result
        self.events.append("narrative")
        if self.fail_narrative:
            raise AIResponseError("Synthetic narrative failure")
        return {"sections": [{"id": d["id"], "claims": [{"text": f"{d['metrics'][d['id']+'.shipments']} shipments in this section.",
            "metric_ids": [d["id"]+".shipments"]}]} for d in context["datasets"]]}


def generate(prompt=EXAMPLE, records=None, ai=None, sections=None, previous=None, **options):
    ai = ai or PlannedAI(sections)
    config = replace(settings, report_ai_planner=True, **options)
    records = shipments() if records is None else records
    request = ReportRequest.model_validate({"report_type": "shipment", "prompt": prompt,
        "data": {"orders": records}, "expected_records": len(records), "previous_report": previous, "filters": {
            "date_field": "manifest_datetime", "from_date": "2026-10-01", "to_date": "2026-10-21", "status": "All Statuses"}})
    return asyncio.run(ReportService(ai, config=config).generate_report(request, user_id="synthetic-account")), ai


def ids(report):
    return [s.id for s in report.structured_report.sections]


def test_arbitrary_section_ids_and_order_are_model_selected_not_fixed_template():
    report, ai = generate()
    assert ids(report) == [s["id"] for s in EXAMPLE_SECTIONS]
    assert len(ai.calls) == 2
    context = json.loads(ai.calls[0][-1]["content"])
    assert context["prompt"] == EXAMPLE
    assert "coverage_requirements" not in context and "datasets" not in context
    assert context["source_scope"]["date_field"] == "manifest_datetime"
    assert context["account_scope"] != "synthetic-account"
    assert all(f["availability"] == "unknown_until_fetch" for f in context["available_source_schema"])
    assert "Synthetic Private Name" not in json.dumps(ai.calls)
    assert report.records_analyzed == 8
    pending = next(t for t in report.tables if t.id == "section_returns_waiting")
    assert sum(row[1] for row in pending.values) == 2
    assert report.tables[-1].columns[:2] == ["Destination City", "Destination Pincode"]
    assert "financial" not in ids(report)


CASES = [
    ("Financial-only request", [section("spend_only", metrics=("shipping_charges", "avg_charge", "median_charge", "order_value"))]),
    ("RTO courier and city analysis", [section("returns_by_route", ("courier", "destination_city"), filters=[{"field": "status_category", "op": "eq", "value": "rto"}])]),
    ("Delivered shipment timelines", [section("transit_duration", ("courier",), ("avg_delivery_days", "median_delivery_days", "max_delivery_days"), filters=[{"field": "status_category", "op": "eq", "value": "delivered"}])]),
    ("Buyer-first summary", EXAMPLE_SECTIONS),
    ("Status chart without financial section", [section("outcome_visual", ("status_category",), charts=["bar"])]),
    ("Sirf kharche ki report banao", [section("spend_only", metrics=("shipping_charges", "avg_charge", "median_charge", "order_value"))]),
    ("Pehle buyers phir courier aur city pincode dikhao", [EXAMPLE_SECTIONS[0], EXAMPLE_SECTIONS[2], EXAMPLE_SECTIONS[3]]),
    ("\u0915\u0947\u0935\u0932 \u0938\u094d\u0925\u093f\u0924\u093f \u0915\u093e \u091a\u093e\u0930\u094d\u091f \u0926\u093f\u0916\u093e\u0913", [section("outcome_visual", ("status_category",), charts=["bar"])]),
    ("Zone and payment combinations ranked by average weight", [section("route_payment_mix", ("zone", "payment_mode"), ("orders", "avg_weight"), ranking_metric="avg_weight")]),
]


@pytest.mark.parametrize("prompt, sections", CASES)
def test_different_prompts_accept_different_composable_plans(prompt, sections):
    records = shipments()
    for i, row in enumerate(records):
        row.update(zone="Zone A", payment_mode="COD", weight=1+i)
    report, ai = generate(prompt, records=records, sections=sections)
    assert ids(report) == [s["id"] for s in sections]
    assert [s["id"] for s in json.loads(ai.calls[1][-1]["content"])["approved_ordered_plan"]] == ids(report)
    assert json.loads(ai.calls[0][-1]["content"])["prompt"] == prompt
    assert len(report.structured_report.charts) == sum(len(s.get("charts", [])) for s in sections)
    if "spend_only" not in ids(report):
        assert not any("Shipping Charge" in t.columns for t in report.tables)


def test_zero_subset_is_explicit_without_disappearing():
    records = shipments()
    for row in records:
        row["status"] = "Delivered"
    report, _ = generate(records=records)
    target = next(s for s in report.structured_report.sections if s.id == "returns_waiting")
    assert any("No matching shipments" in b.text for b in target.blocks if b.type == "paragraph")
    assert ids(report) == [s["id"] for s in EXAMPLE_SECTIONS]


@pytest.mark.parametrize("dimensions, metrics", [(("rto_reason",), ("orders",)), (("destination_pincode",), ("orders",)), ((), ("avg_weight",))])
def test_unavailable_fields_preserve_requested_section(dimensions, metrics):
    records = shipments()
    for row in records:
        row["consignee"].pop("pincode", None)
        row["cancellation_reason"] = "Synthetic cancellation reason"
    report, _ = generate("Requested missing field analysis", records=records,
        sections=[section("requested_detail", dimensions, metrics), section("available_volume")])
    assert ids(report) == ["requested_detail", "available_volume"]
    assert "No data available" in report.content.sections[0].content
    assert all(table.id != "section_requested_detail" for table in report.tables)
    assert all(card.value != "Unavailable" for card in report.structured_report.metrics)
    assert "Synthetic cancellation reason" not in report.content.sections[0].content


def test_missing_field_can_request_clarification():
    with pytest.raises(ClarificationRequiredError):
        generate(sections=[section("needed", ("rto_reason",), missing_data="clarify")])


def test_approximate_consignee_groups_are_labeled():
    records = shipments()
    for row in records:
        row["consignee"].pop("id")
    report, _ = generate(records=records)
    assert "approximate" in " ".join(report.structured_report.limitations)


def test_narrative_failure_preserves_all_approved_tables():
    report, _ = generate()
    fallback, _ = generate(ai=PlannedAI(fail_narrative=True))
    assert ids(report) == ids(fallback)
    assert [t.values for t in report.tables] == [t.values for t in fallback.tables]
    assert fallback.generation["narrative_source"] == "deterministic_fallback"


@pytest.mark.parametrize("defect", ["sql", "invalid_cohort"])
def test_invalid_plans_never_fetch_or_publish(defect):
    def alter(plan):
        if defect == "sql": plan["sql"] = "SELECT *"
        else: plan["cohort"] = [{"field": "charge", "op": "between", "value": [10]}]
        return plan
    with pytest.raises(AIResponseError):
        generate(ai=PlannedAI(alter_plan=alter))


@pytest.mark.parametrize("defect, kept, note", [
    ("unknown_metric", ["returns_waiting", "partner_volume", "postal_destinations"], "Profit python is not available"),
    ("unknown_dimension", [s["id"] for s in EXAMPLE_SECTIONS], "Invented dimension is not available"),
    ("missing_requirement", [s["id"] for s in EXAMPLE_SECTIONS[:-1]], None),
    ("duplicate_id", ["recipient_overview", "recipient_overview_2", "partner_volume", "postal_destinations"], None),
    ("invalid_filter", ["returns_waiting", "partner_volume", "postal_destinations"], "its filter could not be applied"),
    ("sku_grouping", [s["id"] for s in EXAMPLE_SECTIONS], "Product name and SKU are not available"),
    ("extra_keys", [s["id"] for s in EXAMPLE_SECTIONS], None),
    ("long_text", [s["id"] for s in EXAMPLE_SECTIONS], None),
])
def test_invalid_plan_parts_are_dropped_and_the_rest_is_reported(defect, kept, note):
    def alter(plan):
        if defect == "unknown_metric": plan["sections"][0]["metrics"] = ["profit_python"]
        elif defect == "unknown_dimension": plan["sections"][0]["dimensions"] = ["invented_dimension"]
        elif defect == "missing_requirement": plan["sections"].pop()
        elif defect == "duplicate_id": plan["sections"][1]["id"] = plan["sections"][0]["id"]
        elif defect == "sku_grouping": plan["sections"][0]["dimensions"] = ["product_name", "sku"]
        elif defect == "extra_keys":
            plan["highlight_rule"] = "red line when freight > 110% of invoice"
            plan["sections"][0]["row_highlight"] = {"color": "red"}
        elif defect == "long_text":
            plan["intent"] = "x" * 400
            plan["sections"][0]["title"] = "y" * 400
        else: plan["sections"][0]["filters"] = [{"field": "charge", "op": "between", "value": [10]}]
        return plan
    ai = PlannedAI(alter_plan=alter)
    report, _ = generate(ai=ai)
    assert ids(report) == kept
    assert ai.events[0] == "plan" and ai.events.count("plan") == 1  # salvaged without a second planning call
    if note:
        assert note in json.dumps(report.model_dump(mode="json"))


def test_requested_flagged_section_is_reworded_not_hidden_or_blocked():
    def flagged(plan):
        plan["sections"][0].update(id="flagged_skus", title="SKUs Flagged for High Freight",
                                   requirement_ids=["req_flagged_skus"])
        plan["requirement_ids"] = [r for s in plan["sections"] for r in s["requirement_ids"]]
        return plan
    report, _ = generate(ai=PlannedAI(alter_plan=flagged))
    assert ids(report) == ["marked_skus"] + [s["id"] for s in EXAMPLE_SECTIONS[1:]]
    assert report.structured_report.sections[0].title == "SKUs Marked for High Freight"


def test_missing_capability_question_is_reported_as_limitation_not_blocking():
    question = "SKU is not in the shipment data, so the SKU breakdown cannot be shown."
    def missing(plan):
        return {**plan, "clarification_kind": "missing_capability", "clarifications": [question]}
    report, _ = generate(ai=PlannedAI(alter_plan=missing))
    assert ids(report) == [s["id"] for s in EXAMPLE_SECTIONS]
    assert question in json.dumps(report.model_dump(mode="json"))


@pytest.mark.parametrize("alter", [
    lambda p: {**p, "clarifications": ["Which ranking metric do you mean?"]},
    lambda p: {**p, "source_scope": {**p["source_scope"], "from_date": "2020-01-01"}},
])
def test_ambiguity_and_source_override_return_structured_clarification(alter):
    with pytest.raises(ClarificationRequiredError) as error:
        generate(ai=PlannedAI(alter_plan=alter))
    assert error.value.details["clarification_required"]


def test_empty_prompt_rejected_before_ai():
    from app.errors import InvalidPromptError
    ai = PlannedAI()
    with pytest.raises(InvalidPromptError):
        generate("", ai=ai)
    assert ai.calls == []


def test_first_call_precedes_data_access_and_aggregations(monkeypatch):
    from app.charts.fields import SHIPMENT_ADAPTER
    events = []
    original = SHIPMENT_ADAPTER.rows
    def fail(data):
        assert events == ["plan"]
        raise RuntimeError("stop after plan")
    object.__setattr__(SHIPMENT_ADAPTER, "rows", fail)
    try:
        with pytest.raises(RuntimeError, match="stop after plan"):
            generate(ai=PlannedAI(events=events))
    finally:
        object.__setattr__(SHIPMENT_ADAPTER, "rows", original)


def test_json_and_exports_share_section_order():
    from app.generators.pdf_generator import generate_pdf
    from app.generators.docx_generator import generate_docx
    report, _ = generate()
    document = Document(BytesIO(generate_docx(report)))
    pdf = "\n".join(p.extract_text() for p in PdfReader(BytesIO(generate_pdf(report))).pages)
    titles = [s.title for s in report.structured_report.sections]
    assert [p.text for p in document.paragraphs if p.style.name == "Heading 1" and p.text in titles] == titles
    assert all(title in pdf for title in titles)


@pytest.mark.parametrize("defect", ["missing", "wrong_dimension", "wrong_count", "wrong_order", "wrong_chart", "stale"])
def test_essential_final_coverage_blocks_mutated_reports(defect):
    from app.analysis.accuracy import DatasetSnapshot
    from app.analysis.plan import AnalysisPlan, ExecutionSection
    from app.analysis.executor import execute_ordered_sections
    from app.charts.fields import get_adapter
    specs = copy.deepcopy(EXAMPLE_SECTIONS)
    specs[2]["charts"] = ["bar"]
    report, _ = generate(sections=specs, report_accuracy_strict=False)
    plan = AnalysisPlan(sections=tuple(ExecutionSection.model_validate(s) for s in report.structured_report.plan["sections"]))
    data = {"orders": shipments()}
    snapshot = DatasetSnapshot.capture(data)
    datasets = execute_ordered_sections(plan, get_adapter("shipment").rows(data))
    if defect == "missing": report.structured_report.sections.pop()
    elif defect == "wrong_dimension":
        block = next(b for b in report.structured_report.sections[-1].blocks if b.type == "table")
        block.column_meta[0].key = "destination_state"
    elif defect == "wrong_count":
        block = next(b for b in report.structured_report.sections[2].blocks if b.type == "table")
        block.values[0][1] = 999
    elif defect == "wrong_order": report.structured_report.sections.reverse()
    elif defect == "wrong_chart": report.structured_report.charts[0].datasets[0].data[0] = 999
    else: report.structured_report.sections.append(copy.deepcopy(report.structured_report.sections[0]))
    with pytest.raises(ReportAccuracyError):
        verify_section_coverage(report, plan, snapshot, datasets)


def test_new_run_does_not_merge_stale_sections():
    old, _ = generate()
    new, ai = generate("Only financial summary", sections=[section("new_costs", metrics=("shipping_charges",))], previous=old.model_dump(mode="json"))
    assert ids(new) == ["new_costs"]
    context = json.loads(ai.calls[0][-1]["content"])
    assert context["follow_up_context"]["section_ids"] == ids(old)
    assert "tables" not in context["follow_up_context"]


@pytest.mark.parametrize("metric", ["delivery_rate", "shipping_charges", "avg_charge", "avg_weight"])
def test_rankings_use_requested_metric(metric):
    records = shipments()
    for i, row in enumerate(records): row["weight"] = 1+i
    report, _ = generate(records=records, sections=[section("ranked", ("courier",), ("orders", metric), ranking_metric=metric)])
    table = report.tables[0]
    values = [r[-1] for r in table.values]
    assert values == sorted(values, reverse=True)


def test_ungrounded_numeric_claim_repaired_without_section_loss():
    class Invented(PlannedAI):
        async def generate_json(self, messages, **kwargs):
            result = await super().generate_json(messages, **kwargs)
            if "[grounded_sections]" in messages[0]["content"]:
                for s in result["sections"]: s["claims"][0]["text"] = "999999 shipments reviewed."
            return result
    report, _ = generate(ai=Invented())
    assert "999999" not in report.model_dump_json()
    assert report.generation["narrative_source"] == "ai_repaired"


def test_production_defaults_to_dynamic_planning(monkeypatch):
    monkeypatch.delenv("REPORT_AI_PLANNER", raising=False)
    assert Settings.from_env().report_ai_planner is True


@pytest.mark.parametrize("cap", [None, 3])
def test_plan_then_complete_backend_fetch_then_narrative(cap):
    import httpx
    from app.services.shipment_source import ShipmentSource
    from tests.test_shipment_source import Upstream, TOKEN
    events = []
    upstream = Upstream(shipments(), cap=cap)
    def transport(request):
        events.append("fetch")
        return upstream(request)
    ai = PlannedAI(events=events)
    cfg = replace(settings, report_ai_planner=True, shipment_api_base_url="https://shipments.test")
    service = ReportService(ai, config=cfg, source=ShipmentSource(cfg, transport=httpx.MockTransport(transport)))
    request = ReportRequest.model_validate({"report_type": "shipment", "source": "shipment_api", "prompt": EXAMPLE,
        "filters": {"date_field": "manifest_datetime", "from_date": "2026-10-01", "to_date": "2026-10-21", "status": "All Statuses"}})
    report = asyncio.run(service.generate_report(request, shipment_credential=TOKEN, user_id="synthetic-account"))
    assert events[0] == "plan" and events[-1] == "narrative"
    assert report.records_analyzed == 8
    assert [(r["page"], r["perPage"]) for r in upstream.requests[:2]] == [(1, 1), (1, 8)]
    assert len(upstream.requests) > 2 if cap else len(upstream.requests) == 2
    assert all(h["authorization"] == TOKEN for h in upstream.headers)
    assert TOKEN not in json.dumps(ai.calls)
    # Prompt-local return filters never narrow the source request or other section denominators.
    assert len({json.dumps({k: v for k, v in r.items() if k not in {"page", "perPage"}}, sort_keys=True) for r in upstream.requests}) == 1


@pytest.mark.parametrize("failure", ["ambiguous", "invalid", "irrelevant"])
def test_plan_rejection_happens_before_any_upstream_request(failure):
    import httpx
    from app.services.shipment_source import ShipmentSource
    from tests.test_shipment_source import Upstream, TOKEN
    upstream = Upstream(shipments())
    def alter(plan):
        if failure in {"ambiguous", "irrelevant"}:
            return {**plan, "clarifications": ["Please specify a shipment analysis."]}
        return {**plan, "python": "run_code()"}
    ai = PlannedAI(alter_plan=alter)
    cfg = replace(settings, report_ai_planner=True, shipment_api_base_url="https://shipments.test")
    request = ReportRequest.model_validate({"report_type": "shipment", "source": "shipment_api", "prompt": "Which is best?",
        "filters": {"date_field": "manifest_datetime", "from_date": "2026-10-01", "to_date": "2026-10-21", "status": "All Statuses"}})
    service = ReportService(ai, config=cfg, source=ShipmentSource(cfg, transport=httpx.MockTransport(upstream)))
    with pytest.raises((ClarificationRequiredError, AIResponseError)):
        asyncio.run(service.generate_report(request, shipment_credential=TOKEN, user_id="synthetic-account"))
    assert upstream.requests == [] and len(ai.calls) == (2 if failure == "invalid" else 1)


def test_dynamic_percentages_use_section_denominator_and_requested_group_dimensions():
    specs = [section("mix", ("courier", "destination_city"), ("orders", "share_percent", "delivery_rate"))]
    report, _ = generate(sections=specs)
    values = report.tables[0].values
    assert sum(row[2] for row in values) == 8
    assert sum(row[3] for row in values) == 100
    assert sum(report.structured_report.metrics[0].value for _ in [0]) == 8


def test_requested_chart_form_without_compatible_fields_is_explicit():
    report, _ = generate(sections=[section("counts", metrics=("orders",), charts=["histogram"])])
    assert report.structured_report.chart_notes[0].requested_type == "histogram"
    assert not report.structured_report.charts


def test_empty_global_dataset_never_generates_narrative_or_report():
    from app.errors import NoMatchingRecordsError
    ai = PlannedAI()
    with pytest.raises(NoMatchingRecordsError):
        generate(records=[], ai=ai)
    assert len(ai.calls) == 1


@pytest.mark.parametrize("kind, metrics", [("histogram", ("shipping_charges",)), ("scatter", ("shipping_charges", "order_value"))])
def test_numeric_charts_reuse_registered_builders_and_current_subset(kind, metrics):
    report, _ = generate(sections=[section("numeric_view", metrics=metrics, charts=[kind])])
    chart = report.structured_report.charts[0]
    assert chart.type == kind
    if kind == "histogram":
        assert sum(chart.datasets[0].data) == 8
    else:
        assert len(chart.datasets[0].points) == 8


def test_report_title_is_short_and_does_not_concatenate_sections():
    specs = [section("counts", ("payment_mode",)),
        section("delivered", ("payment_mode",), ("delivery_rate",)),
        section("returned", ("payment_mode",), ("rto_rate",))]
    records = shipments()
    for row in records: row["payment_mode"] = "COD"
    ai = PlannedAI(specs, alter_plan=lambda p: {**p, "report_title": "COD vs Prepaid Performance"})
    report, _ = generate(records=records, ai=ai)
    assert report.content.title == "COD vs Prepaid Performance"
    assert report.structured_report.title == report.report_name == report.content.title
    assert len(report.structured_report.sections) == 3


def test_missing_report_title_uses_shared_subject_instead_of_all_headings():
    from app.analysis.interpreter import concise_report_title
    from app.analysis.plan import AnalysisPlan, ExecutionSection
    from app.charts.fields import get_adapter
    specs = [section("counts", ("payment_mode",)), section("rates", ("payment_mode",), ("delivery_rate",))]
    plan = AnalysisPlan(sections=tuple(ExecutionSection.model_validate(s) for s in specs))
    assert concise_report_title(plan, get_adapter("shipment")) == "Payment Mode Analysis"



@pytest.mark.parametrize("clarification", [False, True])
def test_general_report_and_graph_generates_without_technical_plan_error(clarification):
    def invalid(plan):
        return {**plan, "clarifications": ["Which metrics?"]} if clarification else {"unexpected": "invalid plan"}
    ai = PlannedAI(alter_plan=invalid)
    report, _ = generate("create a report and add a graph with it", ai=ai)
    assert report.records_analyzed == 8
    assert report.structured_report.charts[0].type == "bar"
    assert report.content.title == "Shipment Overview"
    assert report.generation["ai_calls"] == (2 if clarification else 3)


def test_invalid_plan_is_repaired_automatically_before_data_fetch():
    attempts = []
    def repair(plan):
        attempts.append(plan)
        return {"unexpected": "bad plan"} if len(attempts) == 1 else plan
    report, ai = generate(ai=PlannedAI(alter_plan=repair))
    assert report.generation["ai_calls"] == 3
    assert ids(report) == [s["id"] for s in EXAMPLE_SECTIONS]
    correction = json.loads(ai.calls[1][-1]["content"])
    assert correction["prompt"] == EXAMPLE and correction["validation_feedback"]


@pytest.mark.parametrize("prompt", [
    "Hey bro can you prepare a simple report with a graph for my shipments?",
    "Mujhe ek report bana do aur graph bhi dikhao",
    "mere data ki report bna do chart ke sath",
    "मुझे एक रिपोर्ट बना दो और ग्राफ भी दिखाओ",
    "मेरे डेटा की रिपोर्ट बनाइए और चार्ट दिखाइए",
])
def test_natural_multilingual_general_requests_survive_bad_planner_format(prompt):
    ai = PlannedAI(alter_plan=lambda p: {"invalid_model_response": True})
    report, _ = generate(prompt, ai=ai)
    assert report.records_analyzed == 8
    assert report.structured_report.charts[0].type == "bar"
    assert report.structured_report.scope["selected_period"]["from_date"] == "2026-10-01"
    assert report.generation["ai_calls"] == 3
    assert json.loads(ai.calls[0][-1]["content"])["prompt"] == prompt


def test_general_request_preserves_explicit_chart_style_in_fallback():
    report, _ = generate("please make a report with a pie chart",
        ai=PlannedAI(alter_plan=lambda p: {"invalid": True}))
    assert report.structured_report.charts[0].type == "pie"


def test_optional_metric_question_does_not_block_specific_analysis_or_replace_sections():
    ai = PlannedAI(alter_plan=lambda p: {**p, "clarification_kind": "optional_preferences",
        "clarifications": ["Which graph style would you prefer?"]})
    report, _ = generate(ai=ai)
    assert ids(report) == [s["id"] for s in EXAMPLE_SECTIONS]
    assert report.generation["ai_calls"] == 2


def test_semantically_general_wording_can_use_model_classification_without_keyword_whitelist():
    def classify(plan):
        return {**plan, "request_kind": "general_analysis", "sections": [], "requirement_ids": []}
    report, _ = generate("Turn these figures into something I can understand, with a visual", ai=PlannedAI(alter_plan=classify))
    assert ids(report) == ["dataset_overview"]
    assert report.structured_report.charts[0].type == "bar"


@pytest.mark.parametrize("prompt", [
    "Show courier delivery performance with a chart",
    "create report for last 90 days",
    "मुझे लाभ की रिपोर्ट बनाओ",
    "Write a romantic report about the moon",
])
def test_general_fallback_never_substitutes_for_specific_or_unrelated_requests(prompt):
    from app.analysis.interpreter import general_report_request
    assert general_report_request(prompt)[0] is False
    with pytest.raises(AIResponseError):
        generate(prompt, ai=PlannedAI(alter_plan=lambda p: {"invalid": True}))


def test_unrelated_intent_returns_a_clear_response_instead_of_fabricated_report():
    from app.errors import InvalidPromptError
    def unrelated(plan):
        return {**plan, "request_kind": "unrelated", "sections": [], "requirement_ids": []}
    with pytest.raises(InvalidPromptError):
        generate("Write a poem", ai=PlannedAI(alter_plan=unrelated))


def test_optional_preferences_cannot_bypass_ui_scope_conflict():
    def conflict(plan):
        return {**plan, "clarification_kind": "optional_preferences", "clarifications": ["Which chart?"],
            "source_scope": {**plan["source_scope"], "from_date": "2020-01-01"}}
    with pytest.raises(ClarificationRequiredError):
        generate("make a report with a graph", ai=PlannedAI(alter_plan=conflict))


def test_general_classification_cannot_drop_a_specific_cohort_and_return_empty_report():
    def empty_scoped_plan(plan):
        return {**plan, "request_kind": "general_analysis", "sections": [], "requirement_ids": [],
            "cohort": [{"field": "status_category", "op": "eq", "value": "delivered"}]}
    with pytest.raises(AIResponseError):
        generate("Show delivered performance", ai=PlannedAI(alter_plan=empty_scoped_plan))


def test_manifest_date_subset_and_operational_statuses_from_screenshot():
    prompt = ("Give me a status breakdown of all orders with status Pending, In Transit, Undelivered, and OFD, "
        "where the manifested date is between 1st Sep 2026 and 10th Oct 2026.")
    records = shipments()
    statuses = ["Pending", "In Transit", "Undelivered", "OFD", "Out For Delivery", "Delivered", "Pending", "Undelivered"]
    manifests = ["2026-09-01", "2026-10-10T18:29:59Z", "2026-10-10T18:30:00Z", "2026-09-15",
        "2026-10-10T10:00:00+05:30", "2026-09-15", "2026-08-31", "2026-09-30"]
    for row, status, manifest in zip(records, statuses, manifests):
        row.update(status=status, manifest_datetime=manifest)
    specs = [section("operational_breakdown", ("shipment_status",), ("orders", "share_percent"), charts=["bar"])]
    cohort = [{"field": "shipment_status", "op": "in", "value": ["Pending", "In Transit", "Undelivered", "OFD"]},
        {"field": "manifest_datetime", "op": "between", "value": ["2026-09-01", "2026-10-10"]}]
    ai = PlannedAI(specs, alter_plan=lambda p: {**p, "cohort": cohort})
    filters = {"date_field": "manifest_datetime", "from_date": "2026-08-24", "to_date": "2026-10-22",
        "status": "All Statuses", "timezone": "Asia/Kolkata"}
    request = ReportRequest.model_validate({"report_type": "shipment", "prompt": prompt, "data": {"orders": records},
        "filters": filters, "expected_records": len(records)})
    report = asyncio.run(ReportService(ai, config=replace(settings, report_ai_planner=True, report_accuracy_strict=True)).generate_report(request))
    assert report.records_analyzed == 5
    assert {r[0]: r[1] for r in report.tables[0].values} == {
        "Pending": 1, "In Transit": 1, "Out For Delivery": 2, "Undelivered": 1}
    assert sum(r[2] for r in report.tables[0].values) == 100
    assert report.structured_report.scope["selected_period"] == filters
    assert len(ai.calls) == 2


@pytest.mark.parametrize("field, op, value", [
    ("manifest_datetime", "between", ["2026-10-10", "2026-09-01"]),
    ("manifest_datetime", "between", ["2026-09-01"]),
    ("manifest_datetime", "gte", "2026-02-30"),
    ("manifest_datetime", "eq", "yesterday"),
])
def test_invalid_date_filters_are_rejected(field, op, value):
    from pydantic import ValidationError
    from app.analysis.plan import ExecutionFilter
    with pytest.raises(ValidationError):
        ExecutionFilter(field=field, op=op, value=value)


def test_prompt_date_filter_outside_ui_requires_clarification_before_fetch():
    ai = PlannedAI(alter_plan=lambda p: {**p, "cohort": [
        {"field": "manifest_datetime", "op": "between", "value": ["2026-09-01", "2026-10-10"]}]})
    with pytest.raises(ClarificationRequiredError):
        generate("Status breakdown from September to October", ai=ai)
    assert len(ai.calls) == 1


def test_missing_manifest_date_never_falls_back_to_order_date():
    from app.charts.fields import get_adapter
    records = shipments()
    for row in records:
        row.pop("manifest_datetime")
    rows = get_adapter("shipment").rows({"orders": records})
    assert all(row["manifest_datetime"] is None and row["order_date"] for row in rows)


@pytest.mark.parametrize("cap", [None, 3])
def test_backend_applies_shared_manifest_filter_upstream_when_rows_omit_it(cap):
    import httpx
    from app.services.shipment_source import ShipmentSource, build_query
    from app.models.report_schema import ReportFilters
    from tests.test_shipment_source import Upstream, TOKEN
    records = shipments()
    for row in records:
        row.pop("manifest_datetime")
    upstream = Upstream(records, cap=cap)
    cohort = [{"field": "manifest_datetime", "op": "between", "value": ["2026-09-01", "2026-10-10"]},
        {"field": "shipment_status", "op": "in", "value": ["Delivered", "Undelivered"]}]
    ai = PlannedAI([section("statuses", ("shipment_status",), required_fields=["shipment_status", "manifest_datetime"],
        filters=[cohort[0]], charts=["bar", "pie"])], alter_plan=lambda p: {**p, "cohort": cohort})
    cfg = replace(settings, report_ai_planner=True, report_accuracy_strict=True, shipment_api_base_url="https://shipments.test")
    filters = {"date_field": "manifest_datetime", "from_date": "2026-08-24", "to_date": "2026-10-22", "status": "All Statuses", "timezone": "Asia/Kolkata"}
    request = ReportRequest.model_validate({"report_type": "shipment", "source": "shipment_api", "prompt": "Status breakdown for September 1 to October 10", "filters": filters})
    service = ReportService(ai, config=cfg, source=ShipmentSource(cfg, transport=httpx.MockTransport(upstream)))
    report = asyncio.run(service.generate_report(request, shipment_credential=TOKEN, user_id="synthetic-account"))
    expected = build_query(ReportFilters.model_validate({**filters, "from_date": "2026-09-01", "to_date": "2026-10-10"}), cfg)
    assert all(r["fromDate"] == expected.from_date and r["toDate"] == expected.to_date for r in upstream.requests)
    assert all(r["type"] == 0 for r in upstream.requests)  # Status subset stays local.
    assert report.records_analyzed == 2
    assert len(report.tables) == 1 and len(report.structured_report.charts) == 2
    assert report.structured_report.chart_notes == []
    assert all(card.value != "Unavailable" for card in report.structured_report.metrics)
    assert report.structured_report.scope["selected_period"]["from_date"] == "2026-09-01"
    assert report.structured_report.scope["authorized_source_scope"] == filters
    assert report.structured_report.scope["source_analysis_filters"][0]["applied_by"] == "shipment_api"
    assert len(upstream.requests) > 2 if cap else len(upstream.requests) == 2


def test_no_usable_analysis_returns_no_data_instead_of_an_unavailable_report():
    from app.errors import NoMatchingRecordsError
    records = shipments()
    for row in records:
        row.pop("manifest_datetime")
    ai = PlannedAI([section("requested", ("shipment_status",))], alter_plan=lambda p: {**p, "cohort": [
        {"field": "manifest_datetime", "op": "between", "value": ["2026-10-01", "2026-10-10"]}]})
    with pytest.raises(NoMatchingRecordsError, match="No data available"):
        generate(records=records, ai=ai, report_accuracy_strict=True)
    assert len(ai.calls) == 1  # No narrative or report is generated.


def test_section_local_date_comparisons_are_not_pushed_into_shared_source():
    from app.analysis.plan import AnalysisPlan, ExecutionSection
    from app.models.report_schema import ReportFilters
    specs = [section("first", ("shipment_status",), filters=[
        {"field": "manifest_datetime", "op": "between", "value": ["2026-10-01", "2026-10-10"]}]),
        section("second", ("shipment_status",), filters=[
        {"field": "manifest_datetime", "op": "between", "value": ["2026-10-11", "2026-10-21"]}])]
    plan = AnalysisPlan(sections=tuple(ExecutionSection.model_validate(s) for s in specs))
    request = ReportRequest(report_type="shipment", source="shipment_api", prompt="Compare two periods", filters=ReportFilters(
        from_date="2026-10-01", to_date="2026-10-21", date_field="manifest_datetime", status="All Statuses"))
    original, updated, pushed, _ = ReportService(PlannedAI())._push_source_dates(request, plan)
    assert original is request and updated is plan and pushed == []


@pytest.mark.parametrize("group_by_date", [False, True])
def test_source_applied_date_requirement_is_removed_only_when_row_values_are_not_used(group_by_date):
    from app.analysis.plan import AnalysisPlan, ExecutionSection
    specs = [section("breakdown", ("manifest_datetime",) if group_by_date else ("shipment_status",),
        required_fields=["manifest_datetime"])]
    plan = AnalysisPlan(sections=tuple(ExecutionSection.model_validate(s) for s in specs))
    request = ReportRequest.model_validate({"report_type": "shipment", "source": "shipment_api", "prompt": "Status breakdown",
        "filters": {"date_field": "manifest_datetime", "from_date": "2026-10-01", "to_date": "2026-10-21", "status": "All Statuses"}})
    _, updated, _, _ = ReportService(PlannedAI())._push_source_dates(request, plan)
    assert ("manifest_datetime" in updated.sections[0].required_fields) is group_by_date


def test_missing_kpi_is_omitted_and_available_later_metrics_fill_the_cards():
    specs = [section("missing_trend", ("manifest_datetime",)), section("couriers", ("courier",)),
        section("zero_returns", ("courier",), filters=[{"field": "rto_stage", "op": "eq", "value": "pending"}]),
        section("cities", ("destination_city",))]
    records = shipments()
    for row in records:
        row.pop("manifest_datetime")
        row["status"] = "Delivered"
    report, _ = generate(records=records, sections=specs, report_accuracy_strict=True)
    cards = report.structured_report.metrics
    assert len(cards) == 3
    assert all("missing_trend" not in c.key and c.value != "Unavailable" for c in cards)
    assert next(c for c in cards if c.key.startswith("zero_returns")).value == 0
    assert [c.value for c in report.key_metrics] == ["8.0", "0.0", "8.0"]


def test_pdf_and_word_do_not_show_unavailable_metric_cards():
    from app.generators.pdf_generator import generate_pdf
    from app.generators.docx_generator import generate_docx
    records = shipments()
    for row in records: row.pop("manifest_datetime")
    report, _ = generate(records=records, sections=[section("missing_trend", ("manifest_datetime",)), section("couriers", ("courier",))])
    word = " ".join(p.text for p in Document(BytesIO(generate_docx(report))).paragraphs)
    pdf = " ".join(p.extract_text() for p in PdfReader(BytesIO(generate_pdf(report))).pages)
    assert "Unavailable" not in word and "Unavailable" not in pdf
    assert "No data available" in word and "No data available" in pdf
