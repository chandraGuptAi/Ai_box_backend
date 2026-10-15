"""Chart skills engine: prompt planning, Python-computed chart data, AI plans, API and exports."""
from collections import Counter
from datetime import date, timedelta
from io import BytesIO
import json

import pytest
from docx import Document

from app.charts import SKILLS, build_charts, get_adapter, plan_charts
from app.charts.planner import chart_planning_instructions
from app.charts.render import FALLBACK_NOTE, chart_table, render_png
from app.models.report_schema import ReportChart
from tests.test_reports import client, fake_ai, isolated_validation_audit, SAMPLE_REQUEST  # noqa: F401

ADAPTER = get_adapter("shipment")
COURIERS = ["Delhivery", "Xpressbees", "Ekart 0.5 KG", "Blue Dart"]
STATUSES = ["Delivered", "Cancelled", "In Transit", "RTO", "Pending", "NDR", "Delivered"]
STATES = ["Maharashtra", "Gujarat", "Delhi", "Karnataka", "Tamil Nadu"]


def orders(n=60):
    return [{"order_status": STATUSES[i % 7], "courier_name": COURIERS[i % 4], "payment_type": "COD" if i % 3 else "Prepaid",
             "order_date": (date(2025, 1, 1) + timedelta(days=i % 40)).isoformat(), "destination_state": STATES[i % 5],
             "totalCharge": str(50 + i), "weight": 0.5 + (i % 7) * 0.5, "invoice_amount": 500 + 10 * i,
             "customer_mobile": "9811100022"} for i in range(n)]


DATA = {"getOrderList": orders()}


def charts_for(prompt, data=DATA, **kwargs):
    return build_charts("shipment", data, plan_charts(prompt, "shipment", ADAPTER), **kwargs)


def one(prompt, data=DATA):
    result = charts_for(prompt, data)
    assert len(result.charts) == 1, (result.charts, result.notes)
    return result.charts[0]


# --- The four specification examples, with independently computed values ----------------------

def test_courier_wise_cancellations_bar():
    chart = one("Show courier-wise cancellations using a bar graph.")
    expected = Counter(o["courier_name"] for o in DATA["getOrderList"] if o["order_status"] == "Cancelled")
    assert chart.type == "bar" and chart.x_axis.field == "courier" and chart.y_axis.field == "cancelled_orders"
    assert dict(zip(chart.labels, chart.datasets[0].data)) == {c: float(expected.get(c, 0)) for c in COURIERS}
    assert chart.title == "Cancelled Orders by Courier" and chart.source == "rule"


def test_trend_for_last_30_days_is_a_daily_line_anchored_on_latest_date():
    chart = one("Show shipment trends for the last 30 days.")
    days = [date(2025, 1, 1) + timedelta(days=i % 40) for i in range(60)]
    end = max(days)
    window = Counter(d for d in days if d > end - timedelta(days=30))
    assert chart.type == "line" and chart.x_axis.field == "order_date"
    assert chart.labels == [(end - timedelta(days=29 - i)).isoformat() for i in range(30)]
    assert chart.datasets[0].data == [float(window.get(date.fromisoformat(l), 0)) for l in chart.labels]
    assert "Last 30 days" in chart.description and end.isoformat() in chart.description


def test_cod_vs_prepaid_doughnut():
    chart = one("Show COD vs prepaid distribution in a donut chart.")
    counts = Counter(o["payment_type"] for o in DATA["getOrderList"])
    assert chart.type == "doughnut" and dict(zip(chart.labels, chart.datasets[0].data)) == {k: float(v) for k, v in counts.items()}


def test_delivery_performance_heatmap_shares_add_up_per_courier():
    chart = one("Show delivery performance using a heatmap.")
    assert chart.type == "heatmap" and chart.value_format.type == "percent"
    assert chart.x_axis.field == "courier" and chart.y_axis.field == "status_category"
    for i, courier in enumerate(chart.labels):
        assert sum(d.data[i] for d in chart.datasets) == pytest.approx(100, abs=0.3)


def test_heatmap_without_two_dimensions_explains_what_is_missing():
    single = {"getOrderList": [{**o, "courier_name": "Delhivery"} for o in orders()]}
    result = charts_for("Show delivery performance using a heatmap.", single)
    assert not result.charts and "at least two courier" in result.notes[0].reason.lower()


# --- Multiple charts, aliases, unsupported requests ---------------------------------------------

def test_dashboard_with_three_charts_keeps_requested_types_in_order():
    result = charts_for("Create a dashboard with shipment status pie chart, courier comparison bar chart, and daily "
                        "shipment line chart.")
    assert [(c.type, c.id) for c in result.charts] == [("pie", "shipment_status"), ("bar", "courier_orders"),
                                                       ("line", "line_order_date")]
    assert not result.notes


@pytest.mark.parametrize("phrase, kind", [
    ("donut chart of couriers", "doughnut"), ("doughnut graph by courier", "doughnut"), ("column chart by state", "bar"),
    ("horizontal bar chart of couriers", "horizontal_bar"), ("stacked bar chart of status by courier", "stacked_bar"),
    ("clustered column chart of payment mode by courier", "grouped_bar"), ("spider chart of couriers", "radar"),
    ("box-and-whisker plot of shipping charges by courier", "boxplot"), ("heat map of status by courier", "heatmap"),
    ("tree map of states", "treemap"), ("speedometer of delivery rate", "gauge"), ("scatterplot of weight vs charges", "scatter"),
    ("bubble chart of couriers", "bubble"), ("histogram of order value", "histogram"), ("sales funnel", "funnel"),
    ("waterfall chart", "waterfall"), ("india map of orders", "map"), ("area graph of daily orders", "area"),
])
def test_aliases(phrase, kind):
    assert [r.type for r in plan_charts(phrase, "shipment", ADAPTER).requests] == [kind]


def test_every_registered_skill_builds_from_valid_data_and_renders():
    prompts = {"bar": "bar chart by courier", "horizontal_bar": "horizontal bar chart of states",
               "grouped_bar": "grouped bar chart of payment mode by courier", "stacked_bar": "stacked bar chart of status by courier",
               "line": "daily line chart", "area": "weekly area chart", "pie": "pie chart of status", "doughnut": "donut of payment",
               "scatter": "scatter plot of weight vs shipping charge", "bubble": "bubble chart of weight vs charges by courier",
               "histogram": "histogram of shipping charges", "heatmap": "heatmap of status by courier",
               "boxplot": "box plot of shipping charges by courier", "treemap": "treemap of states", "funnel": "funnel chart",
               "radar": "radar chart of couriers", "waterfall": "waterfall chart", "gauge": "gauge of delivery rate",
               "map": "india map of orders"}
    assert set(prompts) == set(SKILLS)
    for kind, prompt in prompts.items():
        chart = one(prompt)
        assert chart.type == kind
        ReportChart.model_validate(chart.model_dump())
        image = render_png(chart)
        assert (image is None) == (not SKILLS[kind].exportable)
        assert image is None or image.startswith(b"\x89PNG")
        assert chart_table(chart)[1]


def test_chart_data_accuracy_for_aggregates():
    rows = DATA["getOrderList"]
    charges = one("bar chart of shipping charges by courier")
    expected = {c: round(sum(50 + i for i, o in enumerate(rows) if o["courier_name"] == c), 2) for c in COURIERS}
    assert dict(zip(charges.labels, charges.datasets[0].data)) == expected
    assert charges.value_format.type == "currency" and charges.value_format.decimals == 2
    hist = one("histogram of shipping charges")
    assert sum(hist.datasets[0].data) == len(rows)
    waterfall = one("waterfall chart")
    data, steps = waterfall.datasets[0].data, waterfall.options["steps"]
    assert data[0] - sum(v for v, s in zip(data[1:-1], steps[1:-1]) if s == "decrease") == data[-1]
    assert data[-1] == sum(1 for o in rows if o["order_status"] == "Delivered")
    funnel = one("funnel chart").datasets[0].data
    assert funnel == sorted(funnel, reverse=True) and funnel[0] == len(rows)
    box = one("box plot of shipping charges by courier")
    ekart = sorted(50 + i for i, o in enumerate(rows) if o["courier_name"] == "Ekart 0.5 KG")
    stats = box.datasets[0].boxes[box.labels.index("Ekart 0.5 KG")]
    assert (stats.min, stats.max, stats.count) == (ekart[0], ekart[-1], len(ekart))
    gauge = one("gauge of delivery rate")
    assert gauge.datasets[0].data == [round(100 * sum(o["order_status"] == "Delivered" for o in rows) / len(rows), 1)]


@pytest.mark.parametrize("prompt, reason", [
    ("pie chart of average shipping charge by courier", "parts of a whole"),
    ("gauge of orders", "single percentage"),
    ("stacked bar chart of delivery rate by courier and payment mode", "count or total"),
])
def test_unsuitable_requests_return_specific_notes(prompt, reason):
    result = charts_for(prompt)
    assert not result.charts and reason in result.notes[0].reason


def test_map_requires_recognisable_states():
    data = {"getOrderList": [{**o, "destination_state": "Zone " + str(i)} for i, o in enumerate(orders())]}
    result = charts_for("india map of orders", data)
    assert not result.charts and "not recognisable" in result.notes[0].reason


def test_unknown_chart_type_is_not_invented():
    assert plan_charts("show a sankey diagram of couriers", "shipment", ADAPTER).requests == []


# --- Missing, empty, null and invalid data -------------------------------------------------------

def test_missing_dates_give_a_note_not_a_chart():
    data = {"getOrderList": [{k: v for k, v in o.items() if k != "order_date"} for o in orders()]}
    result = charts_for("daily orders line chart and pie chart of status", data)
    assert [c.type for c in result.charts] == ["pie"]
    assert result.notes[0].requested_type == "Line chart" and result.notes[0].missing == ["Order Date"]


def test_empty_scope_gives_notes_for_every_chart():
    result = charts_for("bar chart by courier, pie chart of status", status_scope="exception_never_present")
    assert not result.charts and len(result.notes) == 2 and "No records" in result.notes[0].reason


@pytest.mark.parametrize("bad", [None, "", "not money", -5, "NaN", "Infinity", True])
def test_null_and_invalid_numbers_are_excluded_not_zeroed(bad):
    rows = orders(10)
    rows[0]["totalCharge"] = bad
    chart = one("histogram of shipping charges", {"getOrderList": rows})
    assert sum(chart.datasets[0].data) == 9
    total = one("bar chart of shipping charges by courier", {"getOrderList": rows})
    assert sum(total.datasets[0].data) == sum(50 + i for i in range(1, 10))


def test_null_categories_are_labelled_not_dropped():
    rows = orders(8)
    rows[0]["courier_name"] = None
    chart = one("bar chart by courier", {"getOrderList": rows})
    assert "Not specified" in chart.labels and sum(chart.datasets[0].data) == 8


# --- Planning: negation, no unwanted charts, AI plans ------------------------------------------------

def test_no_unwanted_charts():
    for prompt in ("Give a short shipment summary.", "Create a detailed report covering delivery trends and tables.",
                   "remove the pie chart"):
        assert plan_charts(prompt, "shipment", ADAPTER).requests == []
    assert plan_charts("remove the pie chart and add a bar chart of couriers", "shipment", ADAPTER).excluded == ["pie"]


def test_ambiguous_chart_is_planned_in_the_same_call_and_validated():
    plan = plan_charts("Write a shipment report and add a donut chart", "shipment", ADAPTER)
    assert [r.resolved for r in plan.requests] == [False]
    system, user = chart_planning_instructions(plan, ADAPTER)
    assert "chart_requests" in system and "Never include numbers" in system and "<charts_to_plan>" in user
    ai = [{"type": "pie", "dimension": "courier", "measure": "orders", "title": "Courier Share"}]
    chart = build_charts("shipment", DATA, plan, ai_requests=ai).charts[0]
    assert (chart.type, chart.x_axis.field, chart.source, chart.title) == ("doughnut", "courier", "ai_plan", "Courier Share")


@pytest.mark.parametrize("bad", [
    [{"type": "doughnut", "dimension": "customer_mobile", "measure": "__import__('os')"}],
    [{"type": "doughnut", "dimension": "courier", "title": "Delivered 999999 orders"}],
    "not a list", [42],
])
def test_invalid_ai_plans_cannot_add_fields_numbers_or_code(bad):
    plan = plan_charts("Write a shipment report and add a donut chart", "shipment", ADAPTER)
    chart = build_charts("shipment", DATA, plan, ai_requests=bad).charts[0]
    assert chart.x_axis.field in {"status_category", "courier"}
    assert "999999" not in chart.title and chart.labels


# --- API: one-call budget, chart-only prompts, compatibility, exports ---------------------------------

def test_chart_only_prompt_makes_no_ai_call(client, fake_ai):
    fake_ai.calls = 0
    response = client.post("/api/reports/generate", json={**SAMPLE_REQUEST, "prompt": "Show courier-wise cancellations using a bar graph."})
    assert response.status_code == 200, response.text
    report = response.json()
    assert fake_ai.calls == 0 and report["generation"]["ai_calls"] == 0
    charts = report["structured_report"]["charts"]
    assert [c["type"] for c in charts] == ["bar"] and report["tables"] == []
    assert len(report["structured_report"]["executive_summary"]["paragraphs"]) == 1


def test_ambiguous_chart_uses_the_single_ai_call(client, fake_ai):
    import copy
    from tests.test_reports import STRUCTURED_AI
    fake_ai.response = {**copy.deepcopy(STRUCTURED_AI),
                        "chart_requests": [{"type": "doughnut", "dimension": "payment_mode", "measure": "orders"}]}
    fake_ai.calls = 0
    response = client.post("/api/reports/generate", json={**SAMPLE_REQUEST, "prompt": SAMPLE_REQUEST["prompt"] + " Add a donut chart."})
    assert response.status_code == 200, response.text
    assert fake_ai.calls == 1 and "<charts_to_plan>" in fake_ai.messages[-1]["content"]
    chart = response.json()["structured_report"]["charts"][0]
    assert (chart["type"], chart["x_axis"]["field"], chart["source"]) == ("doughnut", "payment_mode", "ai_plan")


def test_prompts_without_chart_types_keep_the_existing_charts(client, fake_ai):
    from app.prompts.registry import get_report_definition
    from app.prompts.report_intent import detect_intent, report_charts, scope_analysis
    prompt = "Create a detailed shipment report covering status, courier performance, cancellations and payment modes with tables and charts."
    response = client.post("/api/reports/generate", json={**SAMPLE_REQUEST, "prompt": prompt})
    intent = detect_intent(prompt)
    analysis = scope_analysis(get_report_definition("shipment").analyze(SAMPLE_REQUEST["data"]), SAMPLE_REQUEST["data"], intent)
    legacy = [c.model_dump() for c in report_charts(analysis, intent)]
    assert response.json()["structured_report"]["charts"] == legacy and legacy
    assert all(c["source"] is None for c in legacy)


def test_unproducible_chart_does_not_fail_the_report_and_exports_show_fallback(client, fake_ai):
    prompt = "Show orders on an india map, a funnel chart, a gauge of orders and a radar chart of couriers."
    response = client.post("/api/reports/generate", json={**SAMPLE_REQUEST, "prompt": prompt})
    assert response.status_code == 200, response.text
    report = response.json()
    structured = report["structured_report"]
    assert [c["type"] for c in structured["charts"]] == ["map", "funnel", "radar"]
    assert structured["chart_notes"][0]["requested_type"] == "Gauge chart"
    assert any("Gauge chart:" in l for l in structured["limitations"])
    headers = {"Authorization": "Bearer " + response.headers["X-Report-Access-Token"]}
    pdf = client.get(f"/api/reports/{report['report_id']}/download?format=pdf", headers=headers)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    docx = client.get(f"/api/reports/{report['report_id']}/download?format=docx", headers=headers)
    text = "\n".join(p.text for p in Document(BytesIO(docx.content)).paragraphs)
    assert FALLBACK_NOTE in text and "Shipment Lifecycle Funnel" in text
    tables = Document(BytesIO(docx.content)).tables
    assert any(t.rows[0].cells[0].text == "State" for t in tables)


def test_exports_do_not_rewrite_engine_chart_types():
    from app.generators.pdf_generator import _resolve_export_charts
    chart = one("bar chart by courier")
    assert _resolve_export_charts([chart], "show a donut chart and a bar chart")[0].type == "bar"


def test_chart_json_has_no_personal_data_and_numbers_stay_numeric():
    result = charts_for("bar chart by courier, histogram of order value, scatter plot of weight vs charges")
    payload = json.dumps([c.model_dump() for c in result.charts])
    assert "9811100022" not in payload
    for chart in result.charts:
        for dataset in chart.datasets:
            assert all(isinstance(v, float) for v in dataset.data)


def test_registry_endpoint_lists_every_skill(client):
    response = client.get("/api/skills/charts/registry")
    assert response.status_code == 200
    skills = {s["type"]: s for s in response.json()["skills"]}
    assert set(skills) == set(SKILLS) and skills["map"]["exportable_as_image"] is False
    assert "courier" in response.json()["report_types"]["shipment"]["dimensions"]
