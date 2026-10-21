"""API tests. DeepSeek is replaced with a fake, so no API key or network is needed."""

from __future__ import annotations

import json
import re
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from app.api.report_routes import get_report_service
from app.config import settings
from app.errors import AIServiceTimeoutError
from app.main import app
from app.services.report_service import ReportService

SAMPLE_REQUEST = json.loads((Path(__file__).parent.parent / "samples" / "shipment_request.json").read_text(encoding="utf-8"))

AI_CONTENT = {
    "title": "Shipment Analysis – January 2025",
    "executive_summary": "12 orders were analysed. 3 (25.0%) were delivered and 4 (33.3%) were cancelled.",
    "sections": [{"heading": "Status Overview", "content": "Delivered: 3 orders.\n\nCancelled: 4 orders."}],
    "findings": [
        {"title": "High cancellation share", "detail": "Cancelled orders are 33.3% of total.", "severity": "critical"},
    ],
    "recommendations": [
        {"action": "Confirm COD orders before dispatch", "rationale": "Addresses cancellations.", "priority": "HIGH"},
    ],
    "data_limitations": "Small dataset of 12 orders.",
}


class FakeAI:
    model = "fake-model"

    def __init__(self, response=None, error: Exception | None = None):
        self.response = AI_CONTENT if response is None else response
        self.error = error
        self.messages: list[dict[str, str]] = []
        self.calls = 0

    async def generate_json(self, messages, **options):
        self.calls += 1
        self.messages = messages
        if self.error:
            raise self.error
        return self.response


@pytest.fixture(autouse=True)
def _strict_numeric_mode():
    """These tests cover the fail-closed path (HTTP 502). The default deterministic fallback
    is covered in tests/test_shipment_insights.py."""
    previous = settings.deterministic_fallback
    object.__setattr__(settings, "deterministic_fallback", False)
    yield
    object.__setattr__(settings, "deterministic_fallback", previous)


@pytest.fixture
def fake_ai():
    return FakeAI()


@pytest.fixture(autouse=True)
def isolated_validation_audit(tmp_path, monkeypatch):
    """Regression runs must not rotate away diagnostics from live requests."""
    import logging
    from app.services import report_service
    audit = logging.getLogger("ai_box.validation_audit")
    original_handlers = list(audit.handlers)
    for handler in original_handlers:
        audit.removeHandler(handler)
    monkeypatch.setattr(report_service, "BASE_DIR", tmp_path)
    yield
    for handler in list(audit.handlers):
        audit.removeHandler(handler)
        handler.close()
    for handler in original_handlers:
        audit.addHandler(handler)


@pytest.fixture
def client(fake_ai):
    app.dependency_overrides[get_report_service] = lambda: ReportService(fake_ai)
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


def test_health(client):
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json()["status"] == "ok"


def test_report_types(client):
    response = client.get("/api/reports/types")
    assert response.status_code == 200
    assert response.json()[0]["report_type"] == "shipment"


def test_generate_shipment_report(client):
    response = client.post("/api/reports/generate", json=SAMPLE_REQUEST)
    assert response.status_code == 200, response.text
    body = response.json()
    metrics = body["metrics"]
    assert body["records_analyzed"] == 12
    assert metrics["total_orders"] == 12
    assert metrics["rates"]["delivered_percent"] == 25.0
    assert metrics["rates"]["cancelled_percent"] == 33.3
    assert metrics["rates"]["rto_percent"] == 8.3
    assert metrics["cancellation_reasons"]["top_reasons"][0] == {
        "reason": "Customer not available", "count": 2, "percent": 50.0,
    }
    # Lenient normalisation of AI output
    assert body["content"]["findings"][0]["severity"] == "high"
    assert body["content"]["recommendations"][0]["priority"] == "high"
    assert body["content"]["data_limitations"] == ["Small dataset of 12 orders."]
    assert body["tables"][0]["title"] == "Shipment Status Summary"


def test_personal_data_is_not_sent_to_ai(client, fake_ai):
    client.post("/api/reports/generate", json=SAMPLE_REQUEST)
    sent = json.dumps(fake_ai.messages)
    for order in SAMPLE_REQUEST["data"]["getOrderList"]:
        for field in ("customer_name", "customer_mobile", "customer_email", "shipping_address", "pincode", "awb_number"):
            assert str(order[field]) not in sent, f"{field} leaked to AI"
    assert "[redacted]" in sent  # phone number inside a cancellation reason


def test_nested_response_shape_is_supported(client):
    payload = {**SAMPLE_REQUEST, "data": {"data": {"getOrderList": SAMPLE_REQUEST["data"]["getOrderList"]}}}
    response = client.post("/api/reports/generate", json=payload)
    assert response.status_code == 200


def test_schema_validation_error(client):
    response = client.post("/api/reports/generate", json={"report_type": "shipment"})
    assert response.status_code == 422
    error = response.json()["error"]
    assert error["code"] == "request_validation_error"
    assert error["category"] == "validation"


def test_missing_order_list(client):
    response = client.post("/api/reports/generate", json={"report_type": "shipment", "data": {"foo": []}})
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "invalid_report_data"


def test_empty_order_list(client):
    response = client.post("/api/reports/generate", json={"report_type": "shipment", "data": {"getOrderList": []}})
    assert response.status_code == 422


def test_unsupported_report_type(client):
    response = client.post("/api/reports/generate", json={"report_type": "finance", "data": {}})
    assert response.status_code == 400
    assert response.json()["error"]["code"] == "unsupported_report_type"


def test_invalid_ai_output_is_generation_error(client, fake_ai):
    fake_ai.response = {"unexpected": True}
    response = client.post("/api/reports/generate", json=SAMPLE_REQUEST)
    assert response.status_code == 502
    error = response.json()["error"]
    assert error["code"] == "ai_invalid_structure"
    assert error["category"] == "report_generation"
    assert "no report was created" in error["message"]
    assert error["details"]["ai_calls"] == 1


def test_ai_timeout(client, fake_ai):
    fake_ai.error = AIServiceTimeoutError("timeout")
    response = client.post("/api/reports/generate", json=SAMPLE_REQUEST)
    assert response.status_code == 504
    assert response.json()["error"]["code"] == "ai_timeout"


def test_request_too_large(client):
    oversized = "x" * (settings.max_request_bytes + 1)
    response = client.post(
        "/api/reports/generate", content=oversized, headers={"Content-Type": "application/json"}
    )
    assert response.status_code == 413
    assert response.json()["error"]["code"] == "request_too_large"


@pytest.mark.parametrize(("path", "magic", "media_type"), [
    ("/api/reports/export/docx", b"PK", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    ("/api/reports/export/pdf", b"%PDF", "application/pdf"),
])
def test_exports(client, path, magic, media_type):
    report = client.post("/api/reports/generate", json=SAMPLE_REQUEST).json()
    response = client.post(path, json=report)
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == media_type
    assert response.content.startswith(magic)
    assert re.fullmatch(r'attachment; filename="Bigship_[A-Za-z0-9-]+_[A-Za-z0-9-]+_\d{4}-\d{2}-\d{2}\.(pdf|docx)"',
                        response.headers["content-disposition"])


def test_export_validation_error(client):
    response = client.post("/api/reports/export/pdf", json={"report_type": "shipment"})
    assert response.status_code == 422


def test_cors_preflight(client):
    response = client.options(
        "/api/reports/generate",
        headers={"Origin": "http://localhost:4200", "Access-Control-Request-Method": "POST"},
    )
    assert response.headers.get("access-control-allow-origin") == "http://localhost:4200"


STRUCTURED_AI = {
    "title": "Shipment Review",
    "executive_summary": {"heading": "Executive Summary", "paragraphs": ["12 supplied shipments were reviewed."], "highlights": ["Review cancellation reasons."]},
    "sections": [{"id": "review", "title": "Requested Analysis", "order": 1,
                  "blocks": [{"type": "paragraph", "text": "Cancellation reasons need review."},
                             {"type": "bullet_list", "items": ["Use supplied records only."]},
                             {"type": "table_reference", "table_id": "table_0"}]}],
    "findings": [], "recommendations": [], "data_limitations": []
}


@pytest.mark.parametrize("prompt", [
    "Generate a complete shipment performance analysis report.",
    "Analyze cancellation reasons and identify couriers with the most cancellations.",
    "Compare courier performance.",
    "Create an executive management report highlighting operational risks.",
    "Review payment mode patterns only.",
])
def test_structured_prompt_contract(client, fake_ai, prompt):
    fake_ai.response = STRUCTURED_AI
    response = client.post("/api/reports/generate", json={**SAMPLE_REQUEST, "prompt": prompt})
    assert response.status_code == 200, response.text
    report = response.json()["structured_report"]
    assert report["schema_version"] == "2.0"
    assert report["metrics"][0]["value"] == 12
    assert len(report["metrics"]) <= 4
    # Selected tables depend on the requested topic; their cells remain backend-authored.
    from app.prompts.report_intent import detect_intent, selected_tables
    from app.prompts.registry import get_report_definition
    expected = selected_tables(get_report_definition("shipment").analyze(SAMPLE_REQUEST["data"]), detect_intent(prompt))
    actual = [b for s in report["sections"] for b in s["blocks"] if b["type"] == "table"]
    assert {b["title"] for b in actual} == {t.title for t in expected.values()}
    for block in actual:
        assert block["rows"] == next(t.rows for t in expected.values() if t.title == block["title"])
    assert prompt in fake_ai.messages[1]["content"]
    assert report["scope"]["denominator"] == "supplied_records"
    for chart in report["charts"]:
        assert len(chart["labels"]) == len(chart["datasets"][0]["data"])


@pytest.mark.parametrize("block", [
    {"type": "paragraph", "text": "<script>alert(1)</script>"},
    {"type": "table_reference", "table_id": "table_999"},
])
def test_reject_invalid_structured_blocks(client, fake_ai, block):
    import copy
    fake_ai.response = copy.deepcopy(STRUCTURED_AI)
    fake_ai.response["sections"][0]["blocks"] = [block]
    assert client.post("/api/reports/generate", json=SAMPLE_REQUEST).status_code == 502


@pytest.mark.parametrize("block", [
    {"type": "html", "text": "Invented html"},
    {"type": "table", "columns": ["Invented"], "rows": [[999999]]},
])
def test_unsupported_blocks_are_dropped_locally_without_a_second_call(client, fake_ai, block):
    import copy
    fake_ai.response = copy.deepcopy(STRUCTURED_AI)
    fake_ai.response["sections"][0]["blocks"].append(block)
    response = client.post("/api/reports/generate", json=SAMPLE_REQUEST)
    assert response.status_code == 200, response.text
    assert fake_ai.calls == 1
    sections = json.dumps(response.json()["structured_report"]["sections"])
    assert "Invented" not in sections and "999999" not in sections


def test_reject_invented_number(client, fake_ai):
    fake_ai.response = {**AI_CONTENT, "executive_summary": "999999 shipments were delivered."}
    assert client.post("/api/reports/generate", json=SAMPLE_REQUEST).status_code == 502


def test_stored_report_access_and_downloads(client, fake_ai):
    fake_ai.response = STRUCTURED_AI
    generated = client.post("/api/reports/generate", json=SAMPLE_REQUEST)
    report = generated.json()
    url = "/api/reports/" + report["report_id"]
    headers = {"Authorization": "Bearer " + generated.headers["X-Report-Access-Token"]}
    assert client.get(url).status_code == 404
    assert client.get(url, headers={"Authorization": "Bearer wrong"}).status_code == 404
    assert client.get(url, headers=headers).json() == report
    fake_ai.error = RuntimeError("No further AI calls allowed")
    for fmt, magic in (("pdf", b"%PDF"), ("docx", b"PK")):
        response = client.get(url + "/download?format=" + fmt, headers=headers)
        assert response.status_code == 200, response.text
        assert response.content.startswith(magic)
    assert client.get(url + "/download?format=html", headers=headers).status_code == 422


def test_paginated_scope_preserves_denominator(client):
    data = {"getOrderList": {"data": SAMPLE_REQUEST["data"]["getOrderList"], "total": 1000}, "delivered": 900}
    response = client.post("/api/reports/generate", json={**SAMPLE_REQUEST, "data": data, "allow_partial": True})
    assert response.status_code == 200
    report = response.json()
    assert report["metrics"]["rates"]["delivered_percent"] == 25.0
    assert report["structured_report"]["scope"]["reported_total"] == 1000


@pytest.mark.parametrize("data", [{"getOrderList": None}, {"getOrderList": [None]}])
def test_null_data(client, data):
    assert client.post("/api/reports/generate", json={**SAMPLE_REQUEST, "data": data}).status_code == 422


def test_exact_status_mapping():
    from app.prompts.shipment_prompt import categorize_status
    assert categorize_status("Out For Delivery") == "in_transit"
    assert categorize_status("RTO Delivered") == "rto"
    assert categorize_status("not really delivered") == "other"
    assert categorize_status("42") == "other"


def test_docx_contains_canonical_blocks(client, fake_ai):
    from io import BytesIO
    from docx import Document
    fake_ai.response = STRUCTURED_AI
    report = client.post("/api/reports/generate", json=SAMPLE_REQUEST).json()
    response = client.post("/api/reports/export/docx", json=report)
    doc = Document(BytesIO(response.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Cancellation reasons need review." in text
    assert "Use supplied records only." in text
    selected = [b for s in report["structured_report"]["sections"] for b in s["blocks"] if b["type"] == "table"]
    assert len(doc.tables) == 1 + len(selected)  # KPIs and only the selected analytical tables


@pytest.mark.parametrize("summary", [
    "12 orders were analysed; 3 delivered (25%) and 4 cancelled (33.30%).",
    "12 orders were analysed; 3 delivered (25.00%) and 4 cancelled (33.3%).",
    "12 orders were analysed.\n\n3 delivered (25%) and 4 cancelled (33.3%).",
])
def test_equivalent_numerical_formats_are_accepted(client, fake_ai, summary):
    fake_ai.response = {**AI_CONTENT, "executive_summary": summary}
    response = client.post("/api/reports/generate", json=SAMPLE_REQUEST)
    assert response.status_code == 200, response.text
    assert response.json()["metrics"]["rates"]["delivered_percent"] == 25.0


def test_structured_numerical_formats_and_metadata(client, fake_ai):
    import copy
    fake_ai.response = copy.deepcopy(STRUCTURED_AI)
    fake_ai.response["executive_summary"]["highlights"] = ["Delivery rate is 25%."]
    fake_ai.response["sections"][0]["id"] = "review_999999"
    fake_ai.response["sections"][0]["blocks"][0]["text"] = "99. Review the 12 supplied orders.\n100. Examine cancellations."
    response = client.post("/api/reports/generate", json=SAMPLE_REQUEST)
    assert response.status_code == 200, response.text


@pytest.mark.parametrize(("text", "metrics"), [
    ("1,000 shipments; delivery rate 25.00%.", {"total": 1000, "rate": 25.0}),
    ("Delivered 12, cancelled 3.", {"total": 12, "cancelled": 3}),
    ("12 orders.\n\n3 delivered.", {"total": 12, "delivered": 3}),
    ("Reference BS10001 is an identifier.", {"total": 12}),
])
def test_numeric_evidence_parsing(text, metrics):
    from app.services.report_evidence import validate_numerical_claims
    validate_numerical_claims({"text": text}, metrics, [])


def test_numeric_keys_and_metadata_do_not_provide_evidence():
    from app.errors import AIResponseError
    from app.services.report_evidence import validate_numerical_claims
    with pytest.raises(AIResponseError):
        validate_numerical_claims({"text": "999999 orders."}, {"metric_999999": 12}, [])


def test_structured_highlight_invented_number_is_removed(client, fake_ai):
    import copy
    fake_ai.response = copy.deepcopy(STRUCTURED_AI)
    fake_ai.response["executive_summary"]["highlights"] = ["999999 orders delivered."]
    response = client.post("/api/reports/generate", json=SAMPLE_REQUEST)
    assert response.status_code == 200, response.text
    assert "999999" not in json.dumps(response.json()["structured_report"])
    assert response.json()["generation"] == {"ai_calls": 1, "narrative_source": "ai_repaired",
        "notice": "Some AI statements were removed because their numbers did not match your data. All numbers shown are verified."}


class SequenceAI(FakeAI):
    def __init__(self, responses):
        super().__init__()
        self.responses = iter(responses)
        self.calls = []

    async def generate_json(self, messages, **options):
        import copy
        self.calls.append(copy.deepcopy(messages))
        return next(self.responses)


def test_numeric_failure_never_makes_a_second_ai_call():
    import asyncio
    from app.models.report_schema import ReportRequest
    bad = {**AI_CONTENT, "executive_summary": "12 supplied orders. 999999 invented shipments. Review records."}
    ai = SequenceAI([bad, STRUCTURED_AI])
    report = asyncio.run(ReportService(ai).generate_report(ReportRequest.model_validate(SAMPLE_REQUEST)))
    assert len(ai.calls) == 1
    assert report.content.executive_summary == "12 supplied orders. Review records."
    assert report.generation["narrative_source"] == "ai_repaired"
    assert report.metrics["total_orders"] == 12


def test_unrepairable_numbers_fail_after_one_call_with_clear_error():
    import asyncio
    from app.errors import AIResponseError
    from app.models.report_schema import ReportRequest
    bad = {**AI_CONTENT, "executive_summary": "999999 invented shipments."}
    ai = SequenceAI([bad, bad, bad])
    with pytest.raises(AIResponseError) as caught:
        asyncio.run(ReportService(ai).generate_report(ReportRequest.model_validate(SAMPLE_REQUEST)))
    assert len(ai.calls) == 1
    assert caught.value.code == "ai_unverified_numbers"
    assert "do not match your data" in caught.value.message
    assert caught.value.details["ai_calls"] == 1
    assert "executive_summary" in caught.value.details["fields"]
    assert "999999" not in json.dumps(caught.value.details)


def test_valid_report_needs_no_correction():
    import asyncio
    from app.models.report_schema import ReportRequest
    ai = SequenceAI([STRUCTURED_AI])
    asyncio.run(ReportService(ai).generate_report(ReportRequest.model_validate(SAMPLE_REQUEST)))
    assert len(ai.calls) == 1


def test_timeout_is_reported_after_one_call():
    import asyncio
    from app.models.report_schema import ReportRequest

    class TimeoutAI(SequenceAI):
        async def generate_json(self, messages, **options):
            self.calls.append(messages)
            raise AIServiceTimeoutError("timeout")

    ai = TimeoutAI([])
    with pytest.raises(AIServiceTimeoutError):
        asyncio.run(ReportService(ai).generate_report(ReportRequest.model_validate(SAMPLE_REQUEST)))
    assert len(ai.calls) == 1


def test_evidence_check_accepts_derived_numbers():
    from app.services.report_evidence import validate_numerical_claims
    metrics = {"total_orders": 3, "status_category_breakdown": [{"category": "delivered", "count": 1}]}
    validate_numerical_claims({"text": "Delivery rate: 33.33%."}, metrics, [])
    validate_numerical_claims({"text": "Delivery rate: 33.3%."}, metrics, [])
    from app.errors import AIResponseError
    with pytest.raises(AIResponseError):
        validate_numerical_claims({"text": "Delivery rate: 33.4%."}, metrics, [])


@pytest.mark.parametrize("text", ["999 shipments.", "99% delivery rate.", "100% delivered."])
def test_observed_numbers_require_typed_evidence(text):
    from app.services.report_evidence import validate_numerical_claims
    from app.errors import AIResponseError
    with pytest.raises(AIResponseError):
        validate_numerical_claims({"text": text}, {"total_orders": 100, "rates": {"delivered_percent": 25}}, [])


@pytest.mark.parametrize("text", ["Target 25% delivery.", "Achieve 90% delivery.", "Resolve delays within 3 days."])
def test_targets_rejected_even_when_number_is_verified(text):
    from app.services.report_evidence import validate_numerical_claims
    from app.errors import AIResponseError
    with pytest.raises(AIResponseError) as error:
        validate_numerical_claims({"recommendations": [{"action": text}]},
                                 {"total_orders": 3, "rates": {"delivered_percent": 25, "cancelled_percent": 90}}, [])
    assert error.value.rejections[0]["semantic_type"] == "proposed_recommendation"


def test_prompt_numbers_are_not_observed_evidence():
    from decimal import Decimal
    from app.services.report_evidence import validate_numerical_claims
    from app.errors import AIResponseError
    with pytest.raises(AIResponseError):
        validate_numerical_claims({"text": "999 orders delivered."}, {"total_orders": 12}, [],
                                 {Decimal(999)}, user_prompt="claim 999 orders")
    validate_numerical_claims({"text": "Top 3 couriers."}, {"total_orders": 12}, [], user_prompt="show top 3 couriers")


def test_verified_dates_identifiers_and_numbering():
    from app.services.report_evidence import validate_numerical_claims
    from app.errors import AIResponseError
    metrics = {"order_date_range": {"from": "2025-01-02", "to": "2025-01-07"}, "total_orders": 12}
    validate_numerical_claims({"text": "Period 2025-01-02 to 2025-01-07. Reference 999999.\n99. Review shipments."}, metrics, [])
    with pytest.raises(AIResponseError) as error:
        validate_numerical_claims({"text": "Period 2024-01-02."}, metrics, [])
    assert error.value.rejections[0]["semantic_type"] == "date"


def test_historical_comparison_is_not_verified_by_matching_number():
    from app.services.report_evidence import validate_numerical_claims
    from app.errors import AIResponseError
    with pytest.raises(AIResponseError):
        validate_numerical_claims({"text": "Last year delivery was 25%."}, {"rates": {"delivered_percent": 25}}, [])


def test_failed_report_is_not_stored_and_has_trace_id(client, fake_ai, caplog):
    from app.services.report_service import _REPORTS
    before = set(_REPORTS)
    fake_ai.response = {**AI_CONTENT, "executive_summary": "999999 shipments."}
    response = client.post("/api/reports/generate", json=SAMPLE_REQUEST)
    assert response.status_code == 502
    error = response.json()["error"]
    assert error["code"] == "ai_unverified_numbers"
    details = error["details"]
    assert len(details["trace_id"]) == 32
    assert details["ai_calls"] == 1 and fake_ai.calls == 1
    assert "999999" not in response.text
    assert "X-Report-Access-Token" not in response.headers
    assert set(_REPORTS) == before
    events = [json.loads(r.message) for r in caplog.records if r.message.startswith('{"event"') and json.loads(r.message)["event"] == "report_numeric_rejection"]
    assert len(events) == 1
    assert all(e["trace_id"] == details["trace_id"] for e in events)
    assert events[0]["value"] == "999999"
    assert events[0]["field"] == "executive_summary"
    assert "verified_metrics" not in events[0]  # detailed evidence stays in the private audit file


@pytest.mark.parametrize("text", ["-12 orders.", "Delivery rate: 33.3%.", "Financial losses were 12."])
def test_wrong_semantic_metric_rejected(text):
    from app.services.report_evidence import validate_numerical_claims
    from app.errors import AIResponseError
    metrics = {"total_orders": 12, "rates": {"delivered_percent": 25, "cancelled_percent": 33.3}}
    with pytest.raises(AIResponseError):
        validate_numerical_claims({"text": text}, metrics, [])


# ---- Regression: production 502 (trace e02ea676…) -------------------------------------------
def _sample_metrics():
    from app.prompts.registry import get_report_definition
    return get_report_definition(SAMPLE_REQUEST["report_type"]).analyze(SAMPLE_REQUEST["data"]).metrics


def _no_order_value_metrics():
    from app.services.report_evidence import derived_percentages
    metrics = {"total_orders": 63, "order_value": None,
               "status_category_breakdown": [{"category": "delivered", "count": 35}, {"category": "cancelled", "count": 13}],
               "field_coverage": {"order_value": {"present": 0}, "courier": {"present": 63}},
               "data_quality_notes": ["No order value field was found in the data."]}
    metrics["derived_percentages"] = derived_percentages(metrics)
    return metrics


@pytest.mark.parametrize("text", [
    "Raw status tables show up to 15 labels; omitted groups are not zero.",
    "Payment modes and destinations show up to 10 groups, and courier tables show up to 200 couriers.",
])
def test_backend_data_quality_note_numbers_are_verified(text):
    from app.services.report_evidence import validate_numerical_claims
    validate_numerical_claims({"data_limitations": [text]}, _sample_metrics(), [])


@pytest.mark.parametrize("text", ["15 orders were delivered.", "15 couriers were used.", "Target 15 labels per table."])
def test_note_numbers_do_not_verify_unrelated_or_target_claims(text):
    from app.services.report_evidence import validate_numerical_claims
    from app.errors import AIResponseError
    with pytest.raises(AIResponseError):
        validate_numerical_claims({"data_limitations": [text]}, _sample_metrics(), [])


@pytest.mark.parametrize("text", [
    "Order value coverage is 0.0%, so financial impact cannot be assessed.",
    "0 of 63 orders have an order value, so revenue impact cannot be calculated.",
    "Cost analysis is unavailable because order value is present for 0 orders.",
])
def test_zero_coverage_in_financial_sentence_is_not_an_amount(text):
    from app.services.report_evidence import validate_numerical_claims
    validate_numerical_claims({"findings": [{"detail": text}]}, _no_order_value_metrics(), [])


@pytest.mark.parametrize("text", ["Financial losses were 0.", "Revenue of ₹0 was recorded.", "Order value coverage is 5.0%, limiting financial analysis."])
def test_invented_amounts_and_wrong_percentages_still_rejected(text):
    from app.services.report_evidence import validate_numerical_claims
    from app.errors import AIResponseError
    with pytest.raises(AIResponseError):
        validate_numerical_claims({"findings": [{"detail": text}]}, _no_order_value_metrics(), [])


def test_summed_category_counts_are_not_verified():
    # 13 cancelled + 2 other = 15 is a model-side sum, not a backend metric.
    from app.services.report_evidence import validate_numerical_claims
    from app.errors import AIResponseError
    with pytest.raises(AIResponseError) as error:
        validate_numerical_claims({"text": "15 orders were cancelled or unmapped."}, _no_order_value_metrics(), [])
    assert error.value.rejections[0]["semantic_type"] == "actual_metric"


def test_rejection_diagnostics_have_context_and_bounded_candidates():
    from app.services.report_evidence import safe_rejections, validate_numerical_claims
    from app.errors import AIResponseError
    with pytest.raises(AIResponseError) as error:
        validate_numerical_claims({"text": "Exactly 999 shipments; call 9876543210 for details."}, _sample_metrics(), [])
    first = safe_rejections(error.value.rejections)[0]
    assert "999" in first["context"] and "9876543210" not in first["context"]
    assert len(first["verified_metrics"]) <= 8 and first["verified_metric_count"] >= len(first["verified_metrics"])


def test_report_quoting_backend_notes_needs_no_correction():
    import asyncio, copy
    from app.models.report_schema import ReportRequest
    content = copy.deepcopy(STRUCTURED_AI)
    content["data_limitations"] = ["Raw status tables show up to 15 labels; omitted groups are not zero."]
    ai = SequenceAI([content])
    report = asyncio.run(ReportService(ai).generate_report(ReportRequest.model_validate(SAMPLE_REQUEST)))
    assert len(ai.calls) == 1
    assert report.structured_report.title == STRUCTURED_AI["title"]


def test_note_unit_ignores_filler_words():
    from app.services.report_evidence import validate_numerical_claims
    from app.errors import AIResponseError
    metrics = {"total_orders": 63, "data_quality_notes": ["Courier name is missing for 50 of 63 orders."]}
    validate_numerical_claims({"text": "Courier name is missing for 50 of 63 orders."}, metrics, [])
    with pytest.raises(AIResponseError):
        validate_numerical_claims({"text": "50 of 63 couriers were late."}, metrics, [])


@pytest.mark.parametrize("action", [
    "Review the 12 supplied shipments.",
    "Next steps: investigate the 4 cancelled orders.",
    "Prioritize the 4 recorded cancellations within the supplied dataset.",
])
def test_recommendations_can_cite_observed_counts(action):
    from app.services.report_evidence import validate_numerical_claims
    validate_numerical_claims({"recommendations": [{"action": action}]},
                             {"total_orders": 12, "cancelled": 4}, [])


def test_observed_count_and_target_in_same_recommendation():
    from app.services.report_evidence import validate_numerical_claims
    from app.errors import AIResponseError
    with pytest.raises(AIResponseError) as caught:
        validate_numerical_claims({"recommendations": [{"action": "Review 12 shipments, target 90% delivery."}]},
                                 {"total_orders": 12, "rates": {"delivered_percent": 90}}, [])
    assert [r["value"] for r in caught.value.rejections] == ["90"]


def test_first_call_repairs_narrative_and_preserves_exports(client, fake_ai, caplog):
    import copy
    fake_ai.response = copy.deepcopy(STRUCTURED_AI)
    fake_ai.response["executive_summary"]["paragraphs"] = [
        "12 supplied shipments were reviewed. Losses were 999999. Review cancellation reasons."
    ]
    generated = client.post("/api/reports/generate", json=SAMPLE_REQUEST)
    assert generated.status_code == 200, generated.text
    assert fake_ai.calls == 1
    report = generated.json()
    assert report["structured_report"]["executive_summary"]["paragraphs"] == [
        "12 supplied shipments were reviewed. Review cancellation reasons."
    ]
    expected_metrics = _sample_metrics()
    expected_metrics.pop("data_quality_notes", None)
    assert report["metrics"] == expected_metrics
    assert report["structured_report"]["sections"][0]["blocks"][2]["rows"][0][1] == "3"
    rejection_events = [json.loads(r.message) for r in caplog.records if r.message.startswith('{"event"')]
    assert sum(e["event"] == "report_numeric_rejection" for e in rejection_events) == 1
    repairs = [e for e in rejection_events if e["event"] == "report_numeric_repair"]
    assert repairs[0]["repair_status"] == "succeeded" and repairs[0]["attempt"] == 1
    headers = {"Authorization": "Bearer " + generated.headers["X-Report-Access-Token"]}
    url = "/api/reports/" + report["report_id"]
    assert client.get(url, headers=headers).json() == report
    assert client.get(url).status_code == 404
    for fmt, magic in (("pdf", b"%PDF"), ("docx", b"PK")):
        response = client.get(url + "/download?format=" + fmt, headers=headers)
        assert response.status_code == 200 and response.content.startswith(magic)


def _rejections_for(narrative):
    from app.services.report_evidence import validate_numerical_claims
    from app.errors import AIResponseError
    try:
        validate_numerical_claims(narrative, _sample_metrics(), [])
    except AIResponseError as exc:
        return exc.rejections
    raise AssertionError("Expected an unsupported numerical claim")


def test_repair_uses_original_offsets_preserving_dates_ids_and_decimals():
    from app.services.report_evidence import repair_unsupported_narrative, validate_numerical_claims
    text = "12 orders had a delivery rate of 25.0%. Period 2025-01-02 to 2025-01-07. Reference BS10001. Losses were 999999. Review records."
    original = {"executive_summary": text}
    rejections = _rejections_for(original)
    repaired = repair_unsupported_narrative(original, rejections)
    assert repaired["executive_summary"] == text.replace(" Losses were 999999.", "")
    assert original["executive_summary"] == text
    validate_numerical_claims(repaired, _sample_metrics(), [])


def test_repair_multiple_exact_claims_in_one_field():
    from app.services.report_evidence import repair_unsupported_narrative
    narrative = {"executive_summary": "12 supplied orders. Losses were 999999. Target 99% delivery. Review records."}
    repaired = repair_unsupported_narrative(narrative, _rejections_for(narrative))
    assert repaired["executive_summary"] == "12 supplied orders. Review records."


@pytest.mark.parametrize("narrative", [
    {"title": "999999 orders", "executive_summary": "Review records."},
    {"executive_summary": "999999 orders."},
    {"executive_summary": {"paragraphs": ["999999 orders."], "highlights": []}},
])
def test_unsafe_repairs_fail_closed(narrative):
    from app.services.report_evidence import NarrativeRepairError, repair_unsupported_narrative
    with pytest.raises(NarrativeRepairError):
        repair_unsupported_narrative(narrative, _rejections_for(narrative))


def test_repair_removes_sentence_even_when_it_also_has_a_verified_number():
    from app.services.report_evidence import repair_unsupported_narrative
    narrative = {"executive_summary": "12 orders caused losses of 999999. Review records."}
    assert repair_unsupported_narrative(narrative, _rejections_for(narrative))["executive_summary"] == "Review records."


def test_repair_drops_items_left_empty_instead_of_failing():
    import copy
    from app.services.report_evidence import repair_unsupported_narrative
    narrative = copy.deepcopy(STRUCTURED_AI)
    narrative["executive_summary"]["highlights"] = ["999999 orders delivered.", "Review records."]
    narrative["findings"] = [{"title": "Losses", "detail": "Losses were 999999.", "severity": "high"},
                             {"title": "999999 lost", "detail": "Review records.", "severity": "info"}]
    narrative["recommendations"] = [{"action": "Review records.", "rationale": "Losses were 999999.", "priority": "high"}]
    repaired = repair_unsupported_narrative(narrative, _rejections_for(narrative))
    assert repaired["executive_summary"]["highlights"] == ["Review records."]
    assert repaired["findings"] == []
    assert repaired["recommendations"] == [{"action": "Review records.", "rationale": "", "priority": "high"}]


def test_stale_or_sanitized_metadata_cannot_repair():
    import copy
    from app.services.report_evidence import NarrativeRepairError, repair_unsupported_narrative
    narrative = {"executive_summary": "12 orders. 999999 orders lost. Review records."}
    rejected = _rejections_for(narrative)
    stale = copy.deepcopy(rejected)
    stale[0]["span"] = {"start": 0, "end": 2}
    with pytest.raises(NarrativeRepairError):
        repair_unsupported_narrative(narrative, stale)
    stale[0].pop("span")
    with pytest.raises(NarrativeRepairError):
        repair_unsupported_narrative(narrative, stale)


def test_repair_happens_on_the_first_response():
    import asyncio
    from app.models.report_schema import ReportRequest
    bad = {**AI_CONTENT, "executive_summary": "12 supplied orders. 999999 orders lost. Review records."}
    ai = SequenceAI([bad, bad, bad])
    report = asyncio.run(ReportService(ai).generate_report(ReportRequest.model_validate(SAMPLE_REQUEST)))
    assert len(ai.calls) == 1
    assert report.content.executive_summary == "12 supplied orders. Review records."


def test_repair_removes_dangling_dependent_statement():
    from app.services.report_evidence import repair_unsupported_narrative
    narrative = {"executive_summary": "12 orders reviewed. Losses were 999999. This demonstrates a major decline. Review records."}
    repaired = repair_unsupported_narrative(narrative, _rejections_for(narrative))
    assert repaired["executive_summary"] == "12 orders reviewed. Review records."


def test_repair_revalidates_instead_of_returning_invalid_report(monkeypatch):
    import asyncio
    from app.models.report_schema import ReportRequest
    from app.errors import AIResponseError
    import app.services.report_service as report_module
    bad = {**AI_CONTENT, "executive_summary": "12 orders. 999999 shipments lost. Review records."}
    monkeypatch.setattr(report_module, "repair_unsupported_narrative", lambda narrative, errors: narrative)
    ai = SequenceAI([bad, bad, bad])
    with pytest.raises(AIResponseError) as error:
        asyncio.run(ReportService(ai).generate_report(ReportRequest.model_validate(SAMPLE_REQUEST)))
    assert len(ai.calls) == 1
    assert error.value.details["ai_calls"] == 1
    assert len(error.value.details["trace_id"]) == 32


FINANCIAL_PROMPT = "Generate a Shipment Financial Analysis Report covering shipping charges, courier and payment-mode comparisons."
FINANCIAL_REQUEST = {"report_type": "shipment", "prompt": FINANCIAL_PROMPT, "data": {"getOrderList": [
    {"order_status": "Delivered", "courier_name": "A", "payment_type": "COD", "totalCharge": "0.10", "collectableAmount": 99999},
    {"order_status": "In Transit", "courier_name": "A", "payment_type": "ToPay", "totalCharge": "0.20", "invoice_amount": 88888},
    {"order_status": "Cancelled", "courier_name": "B", "payment_type": "Prepaid", "totalCharge": "10.00"},
    {"order_status": "Pending", "courier_name": "B", "totalCharge": None},
    {"order_status": "Pending", "courier_name": "B", "totalCharge": "not money"}
]}}


def _financial_ai():
    from app.prompts.registry import get_report_definition
    analysis = get_report_definition("shipment").analyze(FINANCIAL_REQUEST["data"])
    index = next(i for i,t in enumerate(analysis.tables) if t.title == "Shipping Charges by Courier")
    return {"title": "Shipment Financial Analysis", "executive_summary": {"paragraphs": ["Total shipping charges are 10.30. Cancelled-order charges are 10.00; they are not proven losses."]},
            "sections": [{"id": "shipping_charges", "title": "Shipping Charges", "order": 1,
                          "blocks": [{"type": "table_reference", "table_id": f"table_{index}"}]}],
            "findings": [], "recommendations": [], "data_limitations": ["Missing charges are excluded from averages; currency is unspecified."]}


def test_financial_report_metrics_and_protected_exports(client, fake_ai):
    fake_ai.response = _financial_ai()
    generated = client.post("/api/reports/generate", json=FINANCIAL_REQUEST)
    assert generated.status_code == 200, generated.text
    report = generated.json()
    financial = report["metrics"]["shipping_financial"]
    assert financial["total_shipping_charges"] == "10.30"
    assert financial["average_charge_per_order"] == "3.43"
    assert financial["charges_associated_with_cancelled_orders"] == "10.00"
    assert financial["cod_shipping_charges"] == "0.10"
    assert financial["topay_shipping_charges"] == "0.20"
    assert financial["prepaid_shipping_charges"] == "10.00"
    assert financial["orders_with_valid_charge"] == 3
    assert financial["orders_with_missing_charge"] == 1
    assert financial["orders_with_invalid_charge"] == 1
    assert financial["courier_charges"][0]["average_charge_per_order"] == "0.15"
    metrics = report["structured_report"]["metrics"]
    assert [m["value"] for m in metrics] == ["10.30", "3.43", "10.00", 5]
    assert len(metrics) == 4
    assert "99999" not in json.dumps(financial) and "88888" not in json.dumps(financial)
    assert "Exact AI JSON schema" in fake_ai.messages[0]["content"]
    headers = {"Authorization": "Bearer " + generated.headers["X-Report-Access-Token"]}
    url = "/api/reports/" + report["report_id"]
    assert client.get(url, headers=headers).json() == report
    fake_ai.error = RuntimeError("No AI calls permitted during downloads")
    for fmt, magic in (("pdf", b"%PDF"), ("docx", b"PK")):
        response = client.get(url + "/download?format=" + fmt, headers=headers)
        assert response.status_code == 200 and response.content.startswith(magic)


@pytest.mark.parametrize("charge", [None, "", "bad", -1, True, "NaN", "Infinity"])
def test_missing_and_invalid_charges_are_not_zero(charge):
    from app.prompts.registry import get_report_definition
    analysis = get_report_definition("shipment").analyze({"getOrderList": [{"totalCharge": charge, "order_status": "Cancelled"}]})
    financial = analysis.metrics["shipping_financial"]
    assert financial["orders_with_valid_charge"] == 0
    assert financial["total_shipping_charges"] is None
    assert financial["average_charge_per_order"] is None
    assert financial["charges_associated_with_cancelled_orders"] is None


def test_invalid_financial_table_block_is_dropped_without_a_second_call(caplog):
    import asyncio, copy
    from app.models.report_schema import ReportRequest
    good = _financial_ai()
    bad = copy.deepcopy(good)
    bad["sections"][0]["blocks"].append({"type": "table", "columns": ["Charge"], "rows": [["10.30"]]})
    ai = SequenceAI([bad, good])
    report = asyncio.run(ReportService(ai).generate_report(ReportRequest.model_validate(FINANCIAL_REQUEST)))
    assert len(ai.calls) == 1
    event = next(json.loads(r.message) for r in caplog.records if r.message.startswith('{"event": "report_schema_rejection"'))
    assert event["location"] == ["sections", 0, "blocks", 1]
    assert event["type"] == "union_tag_invalid"
    assert event["message"]
    assert report.structured_report.metrics[0].value == "10.30"


@pytest.mark.parametrize("bad", [
    {"executive_summary": {"paragraphs": ["Review records."]}},
    {**STRUCTURED_AI, "charts": [{"data": "malformed"}]},
])
def test_schema_problems_are_cleaned_locally_with_one_call(bad):
    import asyncio
    from app.models.report_schema import ReportRequest
    ai = SequenceAI([bad, bad])
    report = asyncio.run(ReportService(ai).generate_report(ReportRequest.model_validate(SAMPLE_REQUEST)))
    assert len(ai.calls) == 1
    assert report.structured_report.executive_summary.paragraphs


def test_unusable_structure_is_a_controlled_failure_after_one_call():
    import asyncio
    from app.models.report_schema import ReportRequest
    from app.errors import AIResponseError
    ai = SequenceAI([{"title": "Only a title"}, STRUCTURED_AI])
    with pytest.raises(AIResponseError) as caught:
        asyncio.run(ReportService(ai).generate_report(ReportRequest.model_validate(SAMPLE_REQUEST)))
    assert len(ai.calls) == 1
    assert caught.value.code == "ai_invalid_structure"
    assert caught.value.details["ai_calls"] == 1
    assert len(caught.value.details["trace_id"]) == 32
    assert "input" not in json.dumps(caught.value.details)


def test_malformed_chart_lengths_remain_rejected():
    from app.models.report_schema import ReportChart
    from pydantic import ValidationError
    with pytest.raises(ValidationError):
        ReportChart(id="bad", title="Bad", labels=["A", "B"], datasets=[{"label": "Orders", "data": [1]}])


def test_null_sections_are_cleaned_locally(client, fake_ai):
    fake_ai.response = {**STRUCTURED_AI, "sections": None}
    response = client.post("/api/reports/generate", json=SAMPLE_REQUEST)
    assert response.status_code == 200, response.text
    assert fake_ai.calls == 1


def test_financial_schema_retry_never_stores_partial_report(client, fake_ai):
    from app.services.report_service import _REPORTS
    before = set(_REPORTS)
    fake_ai.response = {"title": "Financial report", "sections": []}
    response = client.post("/api/reports/generate", json=FINANCIAL_REQUEST)
    assert response.status_code == 502
    assert set(_REPORTS) == before
    assert "X-Report-Access-Token" not in response.headers


# Regression for production trace 7f08cd627a7940b3b566ed46c1d63667.
# This evidence fixture reproduces logged values only; it is never live data.
def _logged_financial_evidence():
    return {"total_orders": 64, "rates": {"cancelled_percent": 23.4},
            "shipping_financial": {"total_shipping_charges": "356037.61",
                "charges_associated_with_cancelled_orders": "28539.00",
                "orders_with_valid_charge": 64, "orders_with_missing_charge": 0, "orders_with_invalid_charge": 0,
                "cancelled_orders_with_valid_charge": 15,
                "courier_charges": [
                    {"orders": 28, "total_shipping_charges": "297747.17", "average_charge_per_order": "10633.83"},
                    {"orders": 17, "total_shipping_charges": "30123.77"},
                    {"orders": 11, "total_shipping_charges": "12932.97"}],
                "payment_mode_charges": [
                    {"payment_mode": "COD", "average_charge_per_order": "20025.54"},
                    {"payment_mode": "Prepaid", "total_shipping_charges": "51356.01", "average_charge_per_order": "1351.47"},
                    {"payment_mode": "ToPay", "total_shipping_charges": "24324.00", "average_charge_per_order": "2027.00"}]},
            "payment_mode_breakdown": [{"payment_mode": "COD", "percent": 21.9}]}


@pytest.mark.parametrize("text", [
    "Charges associated with cancelled orders total 28539.00; these are recorded charges, not verified losses or revenue.",
    "recorded charges on cancelled orders: 28539.00.",
    "Orders with valid charge: 64; missing: 0; invalid: 0.",
    "LTL TCI Express carries the largest share of both orders (28) and charges (297747.17), with an average charge of 10633.83. LTL Delhivery (17 orders, 30123.77) and LTL Ekart (11 orders, 12932.97) follow.",
    "Prepaid totals 51356.01 (average 1351.47) and ToPay totals 24324.00 (average 2027.00).",
    "Charges associated with cancelled orders total 28539.00 across 15 cancelled orders with valid charges.",
    "15 orders are cancelled (23.4% of total), with recorded charges of 28539.00. None of the 15 cancelled orders have a cancellation reason recorded, preventing root-cause analysis.",
    "LTL TCI Express represents 297747.17 of 356037.61 total charges across 28 orders, with an average charge of 10633.83, far above other couriers.",
    "COD orders average 20025.54 per order versus 1351.47 for Prepaid and 2027.00 for ToPay, despite representing only 21.9% of orders.",
])
def test_logged_financial_claims_use_correct_units(text):
    from app.services.report_evidence import validate_numerical_claims
    validate_numerical_claims({"executive_summary": text}, _logged_financial_evidence(), [])


@pytest.mark.parametrize("text", [
    "Orders with valid charge: 28539; missing: 0; invalid: 0.",
    "Charges associated with cancelled orders total 28540.00.",
    "Prepaid totals 51357.01 (average 1351.47).",
    "Orders (356037) carry charges of 356037.61.",
])
def test_financial_unit_fix_does_not_allow_wrong_values(text):
    from app.services.report_evidence import validate_numerical_claims
    from app.errors import AIResponseError
    with pytest.raises(AIResponseError):
        validate_numerical_claims({"executive_summary": text}, _logged_financial_evidence(), [])


def test_logged_financial_response_does_not_need_ai_correction():
    import asyncio, copy
    from types import SimpleNamespace
    content = copy.deepcopy(STRUCTURED_AI)
    content["executive_summary"]["paragraphs"] = [
        "Charges associated with cancelled orders total 28539.00; these are recorded charges, not verified losses or revenue.",
        "Prepaid totals 51356.01 (average 1351.47) and ToPay totals 24324.00 (average 2027.00)."
    ]
    content["sections"][0]["blocks"] = [{"type": "bullet_list", "items": ["Orders with valid charge: 64; missing: 0; invalid: 0."]}]
    ai = SequenceAI([content])
    analysis = SimpleNamespace(metrics=_logged_financial_evidence(), tables=[])
    result = asyncio.run(ReportService(ai)._generate_verified_content([], analysis, FINANCIAL_PROMPT))
    assert result == content
    assert len(ai.calls) == 1


# ---- Regression: financial report 502 (trace 0e0a910f…) -------------------------------------
def _financial_metrics():
    return {"total_orders": 64,
            "status_category_breakdown": [{"category": "delivered", "count": 40}, {"category": "cancelled", "count": 15}],
            "courier_performance": [{"courier": "A", "total_orders": 17, "cancelled": 10}, {"courier": "B", "total_orders": 5, "cancelled": 3},
                                    {"courier": "C", "total_orders": 2, "cancelled": 2}],
            "shipping_financial": {"total_shipping_charges": "5563.09", "average_charge_per_order": "86.92",
                                   "charges_associated_with_cancelled_orders": "1771.99", "orders_with_valid_charge": 64,
                                   "orders_with_missing_charge": 0, "cancelled_orders_with_valid_charge": 15,
                                   "courier_charges": [{"courier": "A", "orders": 17, "total_shipping_charges": "448.94"}]}}


@pytest.mark.parametrize("text", [
    "Across 64 orders, total shipping charges were 5563.09.",
    "Shipping charges on 15 cancelled shipments totalled 1771.99, with an average charge per order of 86.92.",
    "Cancelled shipment charges by courier: A (10), B (3) and C (2); 0 orders had a missing charge.",
    "All 64 orders had a valid shipping charge, and courier A accounted for 448.94 in charges.",
])
def test_financial_sentences_classify_counts_and_amounts_locally(text):
    from app.services.report_evidence import validate_numerical_claims
    validate_numerical_claims({"sections": [{"blocks": [{"text": text}]}]}, _financial_metrics(), [])


@pytest.mark.parametrize("text", [
    "Total shipping charges were 64.",                      # a count presented as money
    "Shipping charges for 64 orders were 5000.00.",         # invented amount
    "The average charge per order was 15.",                 # count presented as money
    "Charges were recorded for 86.92 orders.",              # amount presented as a count
    "Cancelled orders cost 1771.99 last year.",            # historical comparison
    "Courier D (11) had the highest shipping charges.",     # unitless number absent from all metrics
])
def test_financial_sentences_still_reject_unsupported_or_mislabelled_numbers(text):
    from app.services.report_evidence import validate_numerical_claims
    from app.errors import AIResponseError
    with pytest.raises(AIResponseError):
        validate_numerical_claims({"sections": [{"blocks": [{"text": text}]}]}, _financial_metrics(), [])


@pytest.mark.parametrize("text", ["Total shipping charges for 64 orders: 5563.09.", "COD orders (1771.99) versus prepaid."])
def test_decimal_after_order_label_is_an_amount(text):
    from app.services.report_evidence import validate_numerical_claims
    validate_numerical_claims({"text": text}, _financial_metrics(), [])


def test_decimal_after_order_label_must_still_be_verified():
    from app.services.report_evidence import validate_numerical_claims
    from app.errors import AIResponseError
    with pytest.raises(AIResponseError):
        validate_numerical_claims({"text": "Total shipping charges for 64 orders: 5563.10."}, _financial_metrics(), [])


# ---- Regression: order value quoted as money (trace fda1be3c…) -------------------------------
def _order_value_metrics():
    return {"total_orders": 64, "order_value": {"orders_with_value": 64, "total": 1234567.0, "average": 19290.11, "min": 250.0, "max": 98000.0},
            "shipping_financial": {"total_shipping_charges": "5563.09"}}


@pytest.mark.parametrize("text", ["Total order value: 1,234,567.", "Order value totalled 1234567.", "Highest order value: 98,000.",
                                  "Average order value was 19,290.11 across 64 orders."])
def test_order_value_figures_are_amounts(text):
    from app.services.report_evidence import validate_numerical_claims
    validate_numerical_claims({"items": [text]}, _order_value_metrics(), [])


@pytest.mark.parametrize("text", ["Total order value: 1,234,568.", "Order value totalled 64.", "Highest order value: 97,000."])
def test_order_value_figures_must_match_verified_amounts(text):
    from app.services.report_evidence import validate_numerical_claims
    from app.errors import AIResponseError
    with pytest.raises(AIResponseError):
        validate_numerical_claims({"items": [text]}, _order_value_metrics(), [])


# ---- Regression: large money figure labelled by a preceding "orders" word (trace 947af1c9…) ----
def _cod_metrics():
    return {"total_orders": 64, "status_category_breakdown": [{"category": "delivered", "count": 40}],
            "order_value": {"orders_with_value": 64, "total": 2345678.0, "average": 36651.22, "min": 250.0, "max": 98000.0},
            "shipping_financial": {"total_shipping_charges": "5563.09", "cod_shipping_charges": "1771.99"}}


@pytest.mark.parametrize("text", ["COD orders: 2345678 in collected order value.", "Value of delivered orders: 2,345,678.",
                                  "Across 64 orders the collected order value is 2345678."])
def test_number_larger_than_any_count_is_checked_as_amount(text):
    from app.services.report_evidence import validate_numerical_claims
    validate_numerical_claims({"items": [text]}, _cod_metrics(), [])


@pytest.mark.parametrize("text", ["COD orders: 2345679.", "There were 2345678 orders.", "Delivered orders: 65 were recorded as 70 shipments."])
def test_large_or_unverified_numbers_still_rejected(text):
    from app.services.report_evidence import validate_numerical_claims
    from app.errors import AIResponseError
    with pytest.raises(AIResponseError):
        validate_numerical_claims({"items": [text]}, _cod_metrics(), [])


def test_audit_file_keeps_redacted_context(client, fake_ai, tmp_path, monkeypatch):
    import logging
    from app.services import report_service
    fake_ai.response = {**AI_CONTENT, "executive_summary": "999999 shipments. Call 9876543210."}
    audit = logging.getLogger("ai_box.validation_audit")
    for h in list(audit.handlers):
        audit.removeHandler(h)
    monkeypatch.setattr(report_service, "BASE_DIR", tmp_path)
    assert client.post("/api/reports/generate", json=SAMPLE_REQUEST).status_code == 502
    lines = [json.loads(l) for l in (tmp_path / "output" / "report_validation.jsonl").read_text().splitlines()]
    numeric = [l for l in lines if l["event"] == "report_numeric_rejection"]
    assert numeric and "context" in numeric[0] and "9876543210" not in json.dumps(lines)
    for h in list(audit.handlers):
        audit.removeHandler(h); h.close()


# ---- Regression: exact rejected sentences from trace d3baa3ef… -------------------------------
def _trace_d3_metrics():
    from app.services.report_evidence import derived_percentages
    m = {"total_orders": 64, "status_category_breakdown": [{"category": "delivered", "count": 40, "percent": 62.5},
                                                           {"category": "cancelled", "count": 15, "percent": 23.4}],
         "shipping_financial": {"total_shipping_charges": "356037.61", "average_charge_per_order": "5563.09",
                                "orders_with_valid_charge": 64, "orders_with_missing_charge": 0, "orders_with_invalid_charge": 0}}
    m["derived_percentages"] = derived_percentages(m)
    return m


@pytest.mark.parametrize("text", [
    "Total shipping charges across the supplied records are 356037.61, with an average charge per order of 5563.09. "
    "All 64 orders have valid charges (0 missing, 0 invalid), so the average denominator is the full order count.",
    "Cancelled orders total 15, a cancellation rate of 23.4 percent.",
    "The cancelled count is 15 and the percent of cancelled orders is 23.4; 15/64 = 23.4375%.",
    "Cancellation rate: 23.4.",
])
def test_trace_d3_sentences_are_accepted(text):
    from app.services.report_evidence import validate_numerical_claims
    validate_numerical_claims({"findings": [{"detail": text}]}, _trace_d3_metrics(), [])


@pytest.mark.parametrize("text", ["Cancellation rate: 25.0.", "Percent of cancelled orders is 23.9.", "All orders have valid charges (3 missing, 0 invalid)."])
def test_trace_d3_wrong_values_still_rejected(text):
    from app.services.report_evidence import validate_numerical_claims
    from app.errors import AIResponseError
    with pytest.raises(AIResponseError):
        validate_numerical_claims({"findings": [{"detail": text}]}, _trace_d3_metrics(), [])


def _dated_financial_evidence():
    metrics = _trace_d3_metrics()
    metrics["order_date_range"] = {"from": "2026-08-17", "to": "2026-10-13"}
    return metrics


@pytest.mark.parametrize("period", [
    "17 August 2026 to 13 October 2026",
    "August 17, 2026 to October 13, 2026",
    "17 Aug 2026 to 13 Oct. 2026",
    "17th August 2026 to October 13th, 2026",
    "2026-08-17 to 2026-10-13",
])
def test_financial_report_dates_are_verified_as_dates(period):
    from app.services.report_evidence import validate_numerical_claims
    text = ("This report analyses shipping charges across 64 supplied shipment records "
            f"covering the period {period}. Total shipping charges recorded were "
            "356,037.61 (currency as supplied).")
    validate_numerical_claims({"executive_summary": text}, _dated_financial_evidence(), [])


@pytest.mark.parametrize("date_text", [
    "14 October 2026", "October 14, 2026", "13 October 2025",
    "31 February 2026", "February 31, 2026", "2026-02-31",
    "18 August 2026",  # Inside the range is not sufficient boundary evidence.
])
def test_unsupported_dates_rejected_with_full_original_span(date_text):
    from app.errors import AIResponseError
    from app.services.report_evidence import validate_numerical_claims, safe_rejections
    text = f"Period ends {date_text}. Review charge coverage."
    with pytest.raises(AIResponseError) as error:
        validate_numerical_claims({"executive_summary": text}, _dated_financial_evidence(), [])
    rejection, = error.value.rejections
    assert rejection["semantic_type"] == "date"
    assert rejection["value"] == date_text
    assert text[rejection["span"]["start"]:rejection["span"]["end"]] == date_text
    assert rejection["verified_metrics"] == ["2026-08-17", "2026-10-13"]
    assert safe_rejections(error.value.rejections)[0]["value"] == date_text


def test_date_masking_preserves_unsupported_metric_offsets():
    from app.errors import AIResponseError
    from app.services.report_evidence import validate_numerical_claims
    text = "Period ends 13 October 2026. There are 999999 shipments."
    with pytest.raises(AIResponseError) as error:
        validate_numerical_claims({"executive_summary": text}, _dated_financial_evidence(), [])
    rejection, = error.value.rejections
    assert rejection["value"] == "999999"
    assert text[rejection["span"]["start"]:rejection["span"]["end"]] == "999999"


def test_financial_report_written_dates_generate_and_export(client, fake_ai):
    from copy import deepcopy
    request = deepcopy(FINANCIAL_REQUEST)
    for row in request["data"]["getOrderList"]:
        row["order_date"] = "2026-08-17"
    request["data"]["getOrderList"][-1]["order_date"] = "2026-10-13"
    good = _financial_ai()
    text = "Period: 17 August 2026 to 13 October 2026. Total shipping charges are 10.30."
    good["executive_summary"]["paragraphs"] = [text]
    ai = SequenceAI([good])  # Any unnecessary correction call fails this test.
    app.dependency_overrides[get_report_service] = lambda: ReportService(ai)
    response = client.post("/api/reports/generate", json=request)
    assert response.status_code == 200, response.text
    assert len(ai.calls) == 1
    report = response.json()
    assert report["structured_report"]["executive_summary"]["paragraphs"] == [text]
    headers = {"Authorization": "Bearer " + response.headers["X-Report-Access-Token"]}
    url = "/api/reports/" + report["report_id"]
    assert client.get(url, headers=headers).json() == report
    for fmt, magic in (("pdf", b"%PDF"), ("docx", b"PK")):
        exported = client.get(url + "/download?format=" + fmt, headers=headers)
        assert exported.status_code == 200 and exported.content.startswith(magic)


def test_unsupported_date_sentence_is_removed_on_the_first_call(client):
    bad = {**AI_CONTENT, "executive_summary": "12 shipments were reviewed. Period: 2 January 2025 to 8 January 2025."}
    ai = SequenceAI([bad, AI_CONTENT])
    app.dependency_overrides[get_report_service] = lambda: ReportService(ai)
    response = client.post("/api/reports/generate", json=SAMPLE_REQUEST)
    assert response.status_code == 200, response.text
    assert len(ai.calls) == 1
    assert response.json()["content"]["executive_summary"] == "12 shipments were reviewed."


def test_unsupported_date_fails_closed_after_one_call_without_saving(client):
    from app.services.report_service import _REPORTS
    bad = {**AI_CONTENT, "executive_summary": "Period: 2 January 2025 to 8 January 2025."}
    ai = SequenceAI([bad, bad, bad])
    app.dependency_overrides[get_report_service] = lambda: ReportService(ai)
    before = set(_REPORTS)
    response = client.post("/api/reports/generate", json=SAMPLE_REQUEST)
    assert response.status_code == 502
    assert len(ai.calls) == 1
    assert set(_REPORTS) == before
    body = response.json()
    assert body["error"]["code"] == "ai_unverified_numbers"
    assert "trace_id" in json.dumps(body)
    assert "8 January" not in json.dumps(body)


def test_unsupported_written_date_sentence_can_be_repaired():
    from app.errors import AIResponseError
    from app.services.report_evidence import validate_numerical_claims, repair_unsupported_narrative
    narrative = {"executive_summary": "Period ends 14 October 2026. Review charge coverage."}
    with pytest.raises(AIResponseError) as error:
        validate_numerical_claims(narrative, _dated_financial_evidence(), [])
    repaired = repair_unsupported_narrative(narrative, error.value.rejections)
    assert repaired == {"executive_summary": "Review charge coverage."}
    validate_numerical_claims(repaired, _dated_financial_evidence(), [])


def test_written_date_requires_date_evidence_even_when_numbers_match_counts():
    from app.errors import AIResponseError
    from app.services.report_evidence import validate_numerical_claims
    with pytest.raises(AIResponseError) as error:
        validate_numerical_claims({"executive_summary": "Period ends 13 October 2026."},
                                  {"total_orders": 2026, "cancelled": 13}, [])
    rejection, = error.value.rejections
    assert rejection["semantic_type"] == "date"
    assert rejection["verified_metrics"] == []


def _payment_charge_evidence():
    metrics = _logged_financial_evidence()
    metrics["shipping_financial"]["cod_shipping_charges"] = "280357.60"
    metrics["payment_mode_breakdown"] = [
        {"payment_mode": "COD", "count": 14, "percent": 21.9},
        {"payment_mode": "Prepaid", "count": 38, "percent": 59.4},
        {"payment_mode": "ToPay", "count": 12, "percent": 18.8},
    ]
    return metrics


PAYMENT_CHARGE_SENTENCE = (
    "Prepaid is the largest payment mode by order count with 38 orders (59.4%), "
    "but COD carries the largest share of shipping charges at 280,357.60 across "
    "14 orders, an average of 20,025.54 per order."
)


@pytest.mark.parametrize("text", [
    PAYMENT_CHARGE_SENTENCE,
    "COD has the largest share of shipping charges: 280357.60.",
    "COD accounts for a share of charges of 280357.60.",
    "COD has the largest share of charges at 280357.60 INR.",
    "Prepaid has an order share of 59.4 percent.",
    "Prepaid order share: 59.4.",
])
def test_financial_share_word_does_not_override_the_amount_unit(text):
    from app.services.report_evidence import validate_numerical_claims
    validate_numerical_claims({"executive_summary": text}, _payment_charge_evidence(), [])


@pytest.mark.parametrize("text,kind", [
    ("COD has the largest share of shipping charges at 280357.61.", "amount"),
    ("COD has a share of shipping charges at 59.4.", "amount"),
    ("COD has a share of shipping charges at 280357.60%.", "percentage"),
    ("COD has a share of shipping charges at 280357.60 percent.", "percentage"),
    ("COD has an order share of 280357.60.", "percentage"),
])
def test_financial_share_still_requires_matching_evidence_type(text, kind):
    from app.errors import AIResponseError
    from app.services.report_evidence import validate_numerical_claims
    with pytest.raises(AIResponseError) as error:
        validate_numerical_claims({"executive_summary": text}, _payment_charge_evidence(), [])
    rejection, = error.value.rejections
    assert rejection["reason"] == f"No verified {kind} matches at the stated precision"


def test_logged_payment_charge_sentence_does_not_require_correction():
    import asyncio
    from types import SimpleNamespace
    from app.models.report_schema import StructuredNarrative
    good = {"title": "Financial Analysis", "executive_summary": {"paragraphs": [PAYMENT_CHARGE_SENTENCE]}}
    ai = SequenceAI([good])
    analysis = SimpleNamespace(metrics=_payment_charge_evidence(), tables=[])
    result = asyncio.run(ReportService(ai)._generate_verified_content(
        [{"role": "system", "content": "Use verified data."}], analysis, FINANCIAL_PROMPT))
    assert len(ai.calls) == 1
    assert StructuredNarrative.model_validate(result).executive_summary.paragraphs == [PAYMENT_CHARGE_SENTENCE]


# Production rejections (validation log, 2026-10-14/15): digits inside courier names were
# checked as weights, so every call failed and was re-sent up to three times.
COURIER_LABEL_METRICS = {"total_orders": 6, "courier_performance": [
    {"courier": "Delhivery 5KG", "total_orders": 2, "exception": 2},
    {"courier": "Delhivery 1KG", "total_orders": 1, "exception": 1},
    {"courier": "Ekart 0.5 KG", "total_orders": 1, "exception": 1, "delivered": 0},
    {"courier": "Delhivery 0.5KG DropShip", "total_orders": 1, "delivered": 0},
    {"courier": "Delhivery Air 2KG", "total_orders": 1, "delivered": 0},
]}


@pytest.mark.parametrize("text", [
    "Delhivery 5KG reports 2 exceptions, Delhivery 1KG reports 1 exception, and Ekart 0.5 KG reports 1 exception.",
    "Review exception handling for Delhivery 5KG, Delhivery 1KG, and Ekart 0.5 KG.",
    "Delhivery 0.5KG DropShip, Delhivery Air 2KG and Ekart 0.5KG each report 0 delivered orders.",
])
def test_numbers_inside_courier_names_are_not_claims(text):
    from app.services.report_evidence import validate_numerical_claims
    validate_numerical_claims({"findings": [{"detail": text}]}, COURIER_LABEL_METRICS, [])


@pytest.mark.parametrize("text", ["Ekart 0.5 KG reports 7 exceptions.", "Ekart shipped 0.5 kg in total."])
def test_courier_names_do_not_hide_other_numbers(text):
    from app.services.report_evidence import validate_numerical_claims
    from app.errors import AIResponseError
    with pytest.raises(AIResponseError) as error:
        validate_numerical_claims({"findings": [{"detail": text}]}, COURIER_LABEL_METRICS, [])
    assert error.value.code == "ai_unverified_numbers"
    assert "Ekart" in error.value.rejections[0]["context"]  # diagnostics keep the readable name


def test_company_name_and_customer_id_in_report_and_exporters():
    from app.generators.pdf_generator import generate_pdf
    from app.generators.docx_generator import generate_docx
    from app.models.report_schema import ReportDocument, ReportContent
    import datetime

    doc = ReportDocument(
        report_type="shipment_summary",
        report_name="Shipment Performance Report",
        generated_at=datetime.datetime.now(datetime.timezone.utc),
        records_analyzed=42,
        company_name="Indu Enterprises",
        customer_id="CUST-1029",
        content=ReportContent(
            title="Operational Overview",
            executive_summary="Summary text here.",
            sections=[],
            findings=[],
            recommendations=[],
            data_limitations=[],
        ),
    )
    pdf_bytes = generate_pdf(doc)
    assert len(pdf_bytes) > 0

    docx_bytes = generate_docx(doc)
    assert len(docx_bytes) > 0

