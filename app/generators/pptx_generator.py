"""Editable PowerPoint export using the same layout as the in-app slide preview."""
import base64
from io import BytesIO
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt
from app.generators.presentation_layout import build_slides, W, H, INK, LIGHT, WHITE


def generate_pptx(report):
    slides=build_slides(report)
    deck=Presentation()
    deck.slide_width,deck.slide_height=Inches(13.333333), Inches(7.5)
    deck.core_properties.title=slides[0]['title']
    scale=deck.slide_width/W
    def unit(value): return int(value*scale)
    def rgb(value): return RGBColor.from_string(value.lstrip('#'))
    for spec in slides:
        slide=deck.slides.add_slide(deck.slide_layouts[6])
        for item in spec['items']:
            x,y,w,h=[unit(item[key]) for key in ('x','y','w','h')]
            if item['kind']=='rect':
                shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,x,y,w,h)
                shape.fill.solid();shape.fill.fore_color.rgb=rgb(item['fill']);shape.line.fill.background()
            elif item['kind']=='text':
                tf=slide.shapes.add_textbox(x,y,w,h+unit(8)).text_frame
                tf.word_wrap=False
                tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
                for index,line in enumerate(item['lines']):
                    p=tf.paragraphs[0] if index==0 else tf.add_paragraph()
                    p.text=line;p.font.name='Arial';p.font.size=Pt(item['size']*.6)
                    p.font.bold=item['bold'];p.font.color.rgb=rgb(item['color'])
                    p.line_spacing=Pt(item['leading']*.6);p.space_before=Pt(0);p.space_after=Pt(0)
            elif item['kind']=='image':
                slide.shapes.add_picture(BytesIO(base64.b64decode(item['data'])),x,y,w,h)
            elif item['kind']=='table':
                table=slide.shapes.add_table(1+len(item['rows']),len(item['columns']),x,y,w,h).table
                for column,width in zip(table.columns,item['col_widths']): column.width=unit(width)
                table.rows[0].height=unit(item['header_h'])
                all_rows=[dict(lines=item['header_lines'],height=item['header_h'])]+item['rows']
                for r,row in enumerate(all_rows):
                    table.rows[r].height=unit(row['height'])
                    for c,lines in enumerate(row['lines']):
                        cell=table.cell(r,c)
                        cell.margin_left=cell.margin_right=unit(16)
                        cell.margin_top=unit(12 if r==0 else 10);cell.margin_bottom=unit(8)
                        cell.fill.solid();cell.fill.fore_color.rgb=rgb(INK if r==0 else LIGHT if r%2 else WHITE)
                        cell.text='\n'.join(lines)
                        cell.text_frame.word_wrap=False
                        for p in cell.text_frame.paragraphs:
                            p.font.name='Arial';p.font.size=Pt(item['size']*.6);p.font.bold=r==0
                            p.font.color.rgb=rgb(WHITE if r==0 else INK)
                            p.line_spacing=Pt(item['size']*1.28*.6);p.space_before=p.space_after=Pt(0)
        slide.notes_slide.notes_text_frame.text='\n\n'.join(spec['notes'])
    output=BytesIO();deck.save(output)
    return output.getvalue()
