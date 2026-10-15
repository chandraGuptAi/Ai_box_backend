"""Word (.docx) report generation with python-docx."""

from __future__ import annotations

from io import BytesIO
from types import SimpleNamespace

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.models.report_schema import ReportDocument, ReportTable
from app.charts.render import FALLBACK_NOTE, chart_table, render_png

PRIMARY = RGBColor(0x1F, 0x3A, 0x5F)
MUTED = RGBColor(0x5F, 0x6B, 0x7A)
HEADER_FILL = "1F3A5F"
ZEBRA_FILL = "F2F5F9"
SEVERITY_COLORS = {
    "high": RGBColor(0xB4, 0x23, 0x18),
    "medium": RGBColor(0xB2, 0x6B, 0x00),
    "low": RGBColor(0x2E, 0x7D, 0x32),
    "info": RGBColor(0x1F, 0x5F, 0x9F),
}


def generate_docx(report: ReportDocument) -> bytes:
    doc = Document()
    _configure(doc)
    company = (report.company_name or (report.structured_report.company_name if report.structured_report else None) or "").strip()
    footer_label = f"{company} · {report.report_name}" if company else report.report_name
    _add_footer(doc, footer_label)

    content = report.content
    structured = report.structured_report
    if structured:
        content = content.model_copy(update={"title": structured.title, "findings": structured.findings,
                                              "recommendations": structured.recommendations})
    _add_title_block(doc, report)

    doc.add_heading(structured.executive_summary.heading if structured else "Executive Summary", level=1)
    _add_paragraphs(doc, "\n\n".join(structured.executive_summary.paragraphs) if structured else content.executive_summary)
    if structured:
        for item in structured.executive_summary.highlights:
            doc.add_paragraph(item, style="List Bullet")

    metrics = [SimpleNamespace(label=m.label, value=str(m.value) + ("%" if m.unit == "percent" else ""), note=None) for m in structured.metrics] if structured else report.key_metrics
    if metrics:
        doc.add_heading("Key Metrics", level=1)
        has_notes = any(metric.note for metric in metrics)
        columns = ["Metric", "Value", "Notes"] if has_notes else ["Metric", "Value"]
        rows = [
            [m.label, m.value, m.note or ""] if has_notes else [m.label, m.value]
            for m in metrics
        ]
        _add_table(doc, columns, rows, align_numbers=False)

    if structured:
        from app.generators.pdf_generator import _resolve_export_charts
        for chart in _resolve_export_charts(structured.charts, report.prompt):
            doc.add_heading(chart.title, level=1)
            image = render_png(chart)
            if image is None:  # e.g. a geographic map: show the exact data instead of an approximation
                doc.add_paragraph(FALLBACK_NOTE)
                _add_table(doc, *chart_table(chart))
                continue
            doc.add_picture(BytesIO(image), width=Cm(16))
            if chart.description:
                doc.add_paragraph(chart.description)
        for section in structured.sections:
            doc.add_heading(section.title, level=1)
            for block in section.blocks:
                if block.type == "paragraph":
                    _add_paragraphs(doc, block.text)
                elif block.type == "bullet_list":
                    for item in block.items:
                        doc.add_paragraph(item, style="List Bullet")
                else:
                    _add_data_table(doc, block)
    else:
        for section in content.sections:
            doc.add_heading(section.heading, level=1)
            _add_paragraphs(doc, section.content)
        if report.tables:
            doc.add_heading("Data Tables", level=1)
            for table in report.tables:
                _add_data_table(doc, table)

    if content.findings:
        doc.add_heading("Key Findings", level=1)
        for index, finding in enumerate(content.findings, start=1):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.keep_with_next = True
            title = paragraph.add_run(f"{index}. {finding.title}  ")
            title.bold = True
            tag = paragraph.add_run(f"[{finding.severity.upper()}]")
            tag.bold = True
            tag.font.size = Pt(8.5)
            tag.font.color.rgb = SEVERITY_COLORS[finding.severity]
            detail = doc.add_paragraph(finding.detail)
            detail.paragraph_format.left_indent = Cm(0.6)

    if content.recommendations:
        doc.add_heading("Recommendations", level=1)
        for index, recommendation in enumerate(content.recommendations, start=1):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.keep_with_next = bool(recommendation.rationale)
            action = paragraph.add_run(f"{index}. {recommendation.action}  ")
            action.bold = True
            tag = paragraph.add_run(f"[{recommendation.priority.upper()} PRIORITY]")
            tag.bold = True
            tag.font.size = Pt(8.5)
            tag.font.color.rgb = SEVERITY_COLORS[recommendation.priority]
            if recommendation.rationale:
                rationale = doc.add_paragraph(recommendation.rationale)
                rationale.paragraph_format.left_indent = Cm(0.6)

    limitations = structured.limitations if structured else content.data_limitations
    if limitations:
        doc.add_heading("Data Limitations", level=1)
        for limitation in limitations:
            doc.add_paragraph(limitation, style="List Bullet")

    disclaimer = doc.add_paragraph()
    disclaimer.paragraph_format.space_before = Pt(18)
    run = disclaimer.add_run(report.disclaimer)
    run.italic = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = MUTED

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _configure(doc: DocumentType) -> None:
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size in (("Title", 24), ("Heading 1", 15), ("Heading 2", 12)):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = PRIMARY
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(16)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(6)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(10)


def _add_title_block(doc: DocumentType, report: ReportDocument) -> None:
    company = (report.company_name or (report.structured_report.company_name if report.structured_report else None) or "").strip()
    customer_id = (report.customer_id or (report.structured_report.customer_id if report.structured_report else None) or "").strip()

    if company:
        comp_para = doc.add_paragraph()
        comp_para.paragraph_format.space_after = Pt(2)
        comp_run = comp_para.add_run(f"PREPARED FOR: {company.upper()}" + (f" (ID: {customer_id})" if customer_id else ""))
        comp_run.bold = True
        comp_run.font.size = Pt(9.5)
        comp_run.font.color.rgb = PRIMARY

    label = doc.add_paragraph()
    run = label.add_run(report.report_name.upper())
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = MUTED

    doc.add_paragraph(report.content.title, style="Title")

    meta = []
    if company:
        meta.append(("Account", f"{company}" + (f" ({customer_id})" if customer_id else "")))
    meta.extend([
        ("Generated", report.generated_at.strftime("%d %b %Y, %H:%M UTC")),
        ("Records analysed", f"{report.records_analyzed:,}"),
    ])
    if report.prompt and not report.structured_report:
        meta.append(("Analysis focus", report.prompt))
    for key, value in meta:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        key_run = paragraph.add_run(f"{key}: ")
        key_run.bold = True
        key_run.font.size = Pt(9.5)
        key_run.font.color.rgb = MUTED
        value_run = paragraph.add_run(value)
        value_run.font.size = Pt(9.5)
        value_run.font.color.rgb = MUTED
    _add_rule(doc)


def _add_paragraphs(doc: DocumentType, text: str) -> None:
    blocks = [block.strip() for block in text.replace("\r\n", "\n").split("\n\n") if block.strip()]
    for block in blocks:
        lines = [line.strip() for line in block.split("\n") if line.strip()]
        if lines and all(line.startswith(("- ", "* ", "• ")) for line in lines):
            for line in lines:
                doc.add_paragraph(line[2:].strip(), style="List Bullet")
        else:
            doc.add_paragraph(" ".join(lines))


def _add_data_table(doc: DocumentType, table: ReportTable) -> None:
    heading = doc.add_heading(table.title, level=2)
    heading.paragraph_format.keep_with_next = True
    if table.rows:
        _add_table(doc, table.columns, table.rows)
    else:
        doc.add_paragraph("No data available.")
    if getattr(table, "note", None):  # e.g. rows limited for display; never silently truncated
        doc.add_paragraph(table.note)


def _add_table(doc: DocumentType, columns: list[str], rows: list[list[str]], align_numbers: bool = True) -> None:
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for cell, title in zip(table.rows[0].cells, columns):
        cell.text = ""
        run = cell.paragraphs[0].add_run(title)
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade(cell, HEADER_FILL)
    _repeat_header(table.rows[0])

    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for col_index, (cell, value) in enumerate(zip(cells, row)):
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            run.font.size = Pt(9.5)
            if align_numbers and col_index > 0 and _is_numeric(value):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            if row_index % 2 == 1:
                _shade(cell, ZEBRA_FILL)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def _is_numeric(value: str) -> bool:
    stripped = value.replace(",", "").replace(".", "", 1).replace("%", "").strip()
    return stripped.isdigit()


def _shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def _add_rule(doc: DocumentType) -> None:
    paragraph = doc.add_paragraph()
    properties = paragraph._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), HEADER_FILL)
    border.append(bottom)
    properties.append(border)


def _add_footer(doc: DocumentType, report_name: str) -> None:
    paragraph = doc.sections[0].footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(f"AI Box · {report_name} · Confidential · Page ")
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED
    page = paragraph.add_run()
    page.font.size = Pt(8)
    page.font.color.rgb = MUTED
    for tag, attrs, text in (
        ("w:fldChar", {"w:fldCharType": "begin"}, None),
        ("w:instrText", {"xml:space": "preserve"}, "PAGE"),
        ("w:fldChar", {"w:fldCharType": "end"}, None),
    ):
        element = OxmlElement(tag)
        for key, value in attrs.items():
            element.set(qn(key), value)
        if text:
            element.text = text
        page._r.append(element)
